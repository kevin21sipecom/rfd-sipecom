"""Cuestionario para el cliente: Word con las preguntas abiertas del RFD y una columna para responder.

El mismo archivo, devuelto con las respuestas, alimenta el modo «Mejorar RFD».
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ID_PREGUNTA = re.compile(r"^P-\d{2,3}$")
MARINO = RGBColor(0x0B, 0x10, 0x33)
SIN_RESPUESTA = {"", "-", "—", "n/a", "pendiente", "por confirmar", "sin respuesta"}


def _sombrear(celda: Any, color: str) -> None:
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:color"), "auto")
    sombra.set(qn("w:fill"), color)
    celda._tc.get_or_add_tcPr().append(sombra)


def crear_cuestionario(datos: dict[str, Any], ruta: Path) -> Path | None:
    preguntas = datos.get("preguntas_cliente") or []
    if not preguntas:
        return None
    doc = Document()
    seccion = doc.sections[0]
    seccion.left_margin = seccion.right_margin = Cm(2)
    estilo = doc.styles["Normal"]
    estilo.font.name = "Segoe UI"
    estilo.font.size = Pt(10)

    titulo = doc.add_paragraph()
    run = titulo.add_run("Preguntas para el cliente")
    run.bold, run.font.size, run.font.color.rgb = True, Pt(18), MARINO
    sub = doc.add_paragraph()
    sub_run = sub.add_run(f"{datos['codigo']} · {datos['titulo']} · {datos['cliente']} · versión {datos.get('version', '0.1')}")
    sub_run.font.size, sub_run.font.color.rgb = Pt(9.5), RGBColor(0x54, 0x5A, 0x74)
    doc.add_paragraph(
        "Para completar el Requerimiento Funcional Detallado necesitamos su confirmación sobre los puntos siguientes. "
        "Escriba la respuesta en la última columna; puede dejar «Pendiente» en lo que aún no esté definido. "
        "Con este archivo respondido se actualizará el RFD sin rehacerlo."
    )

    tabla = doc.add_table(rows=1, cols=4)
    tabla.style = "Table Grid"
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    anchos = [Cm(1.6), Cm(3.0), Cm(7.4), Cm(5.4)]
    for celda, texto, ancho in zip(tabla.rows[0].cells, ["N.º", "Tema", "Contexto y pregunta", "Respuesta del cliente"], anchos):
        celda.width = ancho
        celda.text = ""
        r = celda.paragraphs[0].add_run(texto)
        r.bold, r.font.color.rgb = True, RGBColor(0xFF, 0xFF, 0xFF)
        _sombrear(celda, "0B1033")
    for p in preguntas:
        fila = tabla.add_row().cells
        for celda, ancho in zip(fila, anchos):
            celda.width = ancho
        fila[0].text = p["id"]
        fila[1].text = p["tema"]
        fila[2].text = ""
        contexto = fila[2].paragraphs[0].add_run(p["contexto"].strip() + " ")
        contexto.font.color.rgb = RGBColor(0x54, 0x5A, 0x74)
        pregunta = fila[2].paragraphs[0].add_run(p["pregunta"].strip())
        pregunta.bold = True
        if p.get("relacion"):
            rel = fila[2].add_paragraph().add_run("Relacionado con: " + ", ".join(p["relacion"]))
            rel.font.size, rel.font.color.rgb = Pt(8), RGBColor(0x8A, 0x90, 0xA8)
        fila[3].text = ""
        _sombrear(fila[3], "F8F9FC")
    ruta.parent.mkdir(parents=True, exist_ok=True)
    doc.save(ruta)
    return ruta


def leer_cuestionario(ruta: Path) -> dict[str, str]:
    """Respuestas de un cuestionario devuelto: {P-01: respuesta}. Ignora las vacías o «Pendiente»."""
    respuestas: dict[str, str] = {}
    for tabla in Document(ruta).tables:
        for fila in tabla.rows:
            celdas = fila.cells
            if len(celdas) < 4:
                continue
            id_ = celdas[0].text.strip()
            if not ID_PREGUNTA.match(id_):
                continue
            texto = "\n".join(p.text.strip() for p in celdas[-1].paragraphs if p.text.strip())
            if texto.strip().lower().rstrip(".") not in SIN_RESPUESTA:
                respuestas[id_] = texto.strip()
    return respuestas

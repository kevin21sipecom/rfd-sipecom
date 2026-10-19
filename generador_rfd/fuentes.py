"""Lectura de transcripciones y contexto con referencias de origen."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

EXTENSIONES = {".txt", ".md", ".docx", ".pptx", ".pdf", ".xlsx"}
MAX_CARACTERES = 120_000
# Claude y los modelos de Codex en la nube admiten contextos mucho mayores que los locales (~4 caracteres por token).
MAX_CARACTERES_CLAUDE = 450_000


def limite_para(proveedor: str) -> int:
    return MAX_CARACTERES_CLAUDE if proveedor in ("claude", "codex") else MAX_CARACTERES


@dataclass(frozen=True)
class Fuente:
    id: str
    nombre: str
    categoria: str
    contenido: str

    def para_prompt(self) -> str:
        return f"[FUENTE {self.id} | {self.categoria} | {self.nombre}]\n{self.contenido}"


def _leer_txt(ruta: Path) -> str:
    raw = ruta.read_bytes()
    try:
        contenido = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        contenido = raw.decode("cp1252")
    return "\n".join(f"L{i}: {line.strip()}" for i, line in enumerate(contenido.splitlines(), 1) if line.strip())


def _leer_docx(ruta: Path) -> str:
    documento = Document(ruta)
    partes: list[str] = []
    p = t = 0
    for bloque in documento.iter_inner_content():
        if isinstance(bloque, Paragraph):
            p += 1
            texto = bloque.text.strip()
            if texto:
                partes.append(f"P{p}: {texto}")
        elif isinstance(bloque, Table):
            t += 1
            for ri, fila in enumerate(bloque.rows, 1):
                for ci, celda in enumerate(fila.cells, 1):
                    texto = " ".join(x.text.strip() for x in celda.paragraphs if x.text.strip())
                    if texto:
                        partes.append(f"T{t}F{ri}C{ci}: {texto}")
    return "\n".join(partes)


def _leer_pptx(ruta: Path) -> str:
    from pptx import Presentation

    partes: list[str] = []
    for d, diapositiva in enumerate(Presentation(ruta).slides, 1):
        for forma in diapositiva.shapes:
            if forma.has_text_frame:
                for parrafo in forma.text_frame.paragraphs:
                    texto = "".join(r.text for r in parrafo.runs).strip()
                    if texto:
                        partes.append(f"D{d}: {texto}")
            if getattr(forma, "has_table", False) and forma.has_table:
                for ri, fila in enumerate(forma.table.rows, 1):
                    for ci, celda in enumerate(fila.cells, 1):
                        if celda.text.strip():
                            partes.append(f"D{d}F{ri}C{ci}: {' '.join(celda.text.split())}")
        if diapositiva.has_notes_slide:
            notas = diapositiva.notes_slide.notes_text_frame.text.strip()
            if notas:
                partes.append(f"D{d}N: {' '.join(notas.split())}")
    return "\n".join(partes)


def _leer_pdf(ruta: Path) -> str:
    """Texto por página (Pg3L12: …). Los PDF escaneados no tienen capa de texto y quedan vacíos."""
    from pypdf import PdfReader

    partes: list[str] = []
    for n, pagina in enumerate(PdfReader(ruta).pages, 1):
        lineas = [x.strip() for x in (pagina.extract_text() or "").splitlines() if x.strip()]
        partes.extend(f"Pg{n}L{i}: {linea}" for i, linea in enumerate(lineas, 1))
    return "\n".join(partes)


def _texto_celda(valor: object) -> str:
    return " ".join(str(getattr(valor, "text", valor)).split())


def _leer_xlsx(ruta: Path) -> str:
    """Una línea por celda, citable como Hoja!A1. Las fórmulas se conservan porque son la lógica de cálculo que
    el programador debe replicar; si Excel guardó su resultado se agrega como [valor: …]."""
    from openpyxl import load_workbook

    partes: list[str] = []
    formulas = valores = None
    try:
        formulas = load_workbook(ruta, read_only=True)
        valores = load_workbook(ruta, read_only=True, data_only=True)
        for hoja in formulas.worksheets:
            hoja.reset_dimensions()  # algunos exportadores escriben una <dimension> errónea y se perderían celdas
            celdas = [(c.coordinate, c.value) for fila in hoja.iter_rows() for c in fila if c.value is not None]
            es_formula = lambda v: _texto_celda(v).startswith("=")  # noqa: E731
            guardados: dict[str, object] = {}
            if any(es_formula(v) for _, v in celdas):
                hoja_valores = valores[hoja.title]
                hoja_valores.reset_dimensions()
                guardados = {c.coordinate: c.value for fila in hoja_valores.iter_rows() for c in fila if c.value is not None}
            for coordenada, valor in celdas:
                texto = _texto_celda(valor)
                if not texto:
                    continue
                if es_formula(valor) and coordenada in guardados:
                    texto += f" [valor: {_texto_celda(guardados[coordenada])}]"
                partes.append(f"{hoja.title}!{coordenada}: {texto}")
    finally:
        for libro in (formulas, valores):
            if libro is not None:
                libro.close()
    return "\n".join(partes)


# Markdown se lee como texto plano: los encabezados, tablas y listas ya son legibles para el modelo.
_LECTORES = {
    ".txt": _leer_txt,
    ".md": _leer_txt,
    ".docx": _leer_docx,
    ".pptx": _leer_pptx,
    ".pdf": _leer_pdf,
    ".xlsx": _leer_xlsx,
}


MAX_ESTILO = 60_000


def leer_documento(ruta: Path) -> str:
    ruta = ruta.expanduser().resolve()
    if not ruta.is_file():
        raise FileNotFoundError(ruta)
    sufijo = ruta.suffix.lower()
    if sufijo == ".ppt":
        raise ValueError(f"{ruta.name} usa el formato antiguo .ppt; ábralo en PowerPoint y guárdelo como .pptx.")
    if sufijo not in EXTENSIONES:
        raise ValueError(f"Formato no admitido: {ruta.name}; use .txt, .md, .docx, .pptx, .pdf o .xlsx.")
    contenido = _LECTORES[sufijo](ruta)
    if not contenido.strip():
        if sufijo == ".pdf":
            raise ValueError(f"{ruta.name} no tiene texto seleccionable (parece escaneado); aplíquele OCR o expórtelo a Word.")
        raise ValueError(f"El archivo no contiene texto extraíble: {ruta.name}")
    return contenido


def cargar_fuentes(
    transcripciones: list[Path],
    contextos: list[Path],
    estilos: list[Path] | None = None,
    limite: int = MAX_CARACTERES,
    avances: list[Path] | None = None,
) -> list[Fuente]:
    """Fuentes de contenido (F01…), RFD en avance a continuar (A01…) y referencias de estilo (E01…), que no aportan hechos."""
    estilos = estilos or []
    avances = avances or []
    if not transcripciones and not contextos and not avances:
        raise ValueError("Se requiere al menos un archivo: una transcripción, un documento de contexto o un RFD en avance.")
    archivos = [("transcripcion", p) for p in transcripciones] + [("contexto", p) for p in contextos]
    nombres = [p.name.lower() for _, p in archivos + [("avance", p) for p in avances] + [("estilo", p) for p in estilos]]
    if len(nombres) != len(set(nombres)):
        raise ValueError("Los archivos deben tener nombres distintos para mantener referencias inequívocas.")
    fuentes = [Fuente(f"F{i:02}", r.name, c, leer_documento(r)) for i, (c, r) in enumerate(archivos, 1)]
    fuentes += [Fuente(f"A{i:02}", r.name, "avance", leer_documento(r)) for i, r in enumerate(avances, 1)]
    total = sum(len(f.contenido) for f in fuentes)
    if total > limite:
        raise ValueError(
            f"Las fuentes contienen {total:,} caracteres; el límite para este motor es {limite:,}. "
            "Quite documentos que no aporten al alcance o divida el RFD."
        )
    for i, ruta in enumerate(estilos, 1):
        fuentes.append(Fuente(f"E{i:02}", ruta.name, "estilo", leer_documento(ruta)[:MAX_ESTILO]))
    return fuentes

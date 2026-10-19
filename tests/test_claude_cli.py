"""Pruebas del generador de RFD.

Las marcadas con `claude` llaman a la CLI real (requieren sesión iniciada y consumen uso).
Omítalas con: pytest -m "not claude". Modelo de prueba: variable RFD_MODELO_TEST (por defecto haiku).
"""

from __future__ import annotations

import json
import os
import shutil
import subprocess
import time
from pathlib import Path
from zipfile import ZipFile

import pytest
from lxml import etree

from generador_rfd import servidor
from generador_rfd.fuentes import Fuente, cargar_fuentes
from generador_rfd.modelo import analizar_fuentes, comando_claude, crear_prompt, temas_de_codigo
from generador_rfd.pipeline import publicar, terminos_ajenos
from generador_rfd.proyecto import inventario_repositorios
from generador_rfd.schema import SCHEMA, auditar_texto_final
from generador_rfd.word import NS, texto_docx

RAIZ = Path(__file__).resolve().parent.parent
EJEMPLOS = RAIZ / "ejemplos"
MODELO = os.environ.get("RFD_MODELO_TEST", "haiku")
CLAUDE = shutil.which("claude")
requiere_claude = pytest.mark.skipif(not CLAUDE, reason="Claude Code CLI no está en PATH")


@pytest.fixture(autouse=True)
def _sin_detalle_real(request, monkeypatch):
    """El detalle por opción llama al modelo: fuera de las pruebas `claude` responde sin especificación."""
    if request.node.get_closest_marker("claude"):
        return
    from generador_rfd import pipeline

    monkeypatch.setattr(pipeline, "detallar_requerimiento",
                        lambda **k: {"vistas": [], "integraciones": [], "secciones_reemplazadas": []})


def test_comando_claude_va_aislado_y_sin_meta_esquema():
    comando = comando_claude("claude", "sonnet", SCHEMA)
    for opcion in ("-p", "--strict-mcp-config", "--no-session-persistence", "--disable-slash-commands"):
        assert opcion in comando
    assert comando[comando.index("--tools") + 1] == ""
    assert comando[comando.index("--setting-sources") + 1] == ""
    esquema = json.loads(comando[comando.index("--json-schema") + 1])
    assert "$schema" not in esquema
    assert esquema["required"] == SCHEMA["required"]


def test_varios_archivos_y_referencias_de_estilo():
    fuentes = cargar_fuentes(
        [EJEMPLOS / "transcripcion_demo.txt"],
        [EJEMPLOS / "contexto_demo.docx", EJEMPLOS / "presentacion_demo.pptx"],
        [EJEMPLOS / "contexto_demo.txt"],
    )
    assert [f.id for f in fuentes] == ["F01", "F02", "F03", "E01"]
    assert [f.categoria for f in fuentes] == ["transcripcion", "contexto", "contexto", "estilo"]
    assert fuentes[2].contenido.startswith("D1:")
    with pytest.raises(ValueError, match="nombres distintos"):
        cargar_fuentes([EJEMPLOS / "transcripcion_demo.txt"], [EJEMPLOS / "transcripcion_demo.txt"])


def test_lee_pdf_y_markdown(tmp_path):
    fuentes = cargar_fuentes([], [EJEMPLOS / "contexto_demo.pdf", EJEMPLOS / "requisitos_demo.md"])
    assert fuentes[0].contenido.startswith("Pg1L1:")
    assert "Proyecto ficticio" in fuentes[0].contenido
    assert "| Solicitudes abiertas por categoría | Semanal | Coordinador |" in fuentes[1].contenido
    vacio = tmp_path / "escaneado.pdf"
    from pypdf import PdfWriter
    w = PdfWriter()
    w.add_blank_page(612, 792)
    w.write(vacio)
    with pytest.raises(ValueError, match="OCR"):
        cargar_fuentes([], [vacio])


def test_lee_excel_por_hoja_y_celda(tmp_path):
    from openpyxl import Workbook
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Requisitos"
    hoja["A1"], hoja["B1"] = "ID", "Descripción"
    hoja["A2"], hoja["B2"] = "R-01", "Alta de solicitudes"
    hoja["C2"] = 3
    libro.create_sheet("Vacía")
    ruta = tmp_path / "requisitos.xlsx"
    libro.save(ruta)
    fuentes = cargar_fuentes([], [ruta])
    assert fuentes[0].contenido.splitlines() == [
        "Requisitos!A1: ID",
        "Requisitos!B1: Descripción",
        "Requisitos!A2: R-01",
        "Requisitos!B2: Alta de solicitudes",
        "Requisitos!C2: 3",
    ]
    assert "Formato no admitido" not in fuentes[0].contenido
    with pytest.raises(ValueError, match="no contiene texto extraíble"):
        cargar_fuentes([], [_libro_vacio(tmp_path / "vacio.xlsx")])


def test_excel_conserva_formulas_como_logica_de_calculo(tmp_path):
    from openpyxl import Workbook
    from generador_rfd.calidad import MARCA_FUENTE
    from generador_rfd.modelo import prompt_analisis
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Calculo"
    hoja["A1"], hoja["B1"], hoja["B2"] = "Total", 2, 3
    hoja["B3"] = "=SUM(B1:B2)"
    ruta = tmp_path / "formulas.xlsx"
    libro.save(ruta)
    # Sin valor guardado (nunca abierto en Excel): la fórmula se conserva igual.
    assert "Calculo!B3: =SUM(B1:B2)" in cargar_fuentes([], [ruta])[0].contenido.splitlines()
    # Con valor guardado: fórmula y resultado juntos.
    hoja_valores = tmp_path / "con_valor.xlsx"
    from zipfile import ZIP_DEFLATED
    with ZipFile(ruta) as entrada, ZipFile(hoja_valores, "w", ZIP_DEFLATED) as salida:
        for item in entrada.infolist():
            datos = entrada.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                datos = datos.replace(b"<v></v>", b"<v>5</v>").replace(b"<v />", b"<v>5</v>")
            salida.writestr(item, datos)
    contenido = cargar_fuentes([], [hoja_valores])[0].contenido
    assert "Calculo!B3: =SUM(B1:B2) [valor: 5]" in contenido.splitlines()
    # El prompt de análisis exige registrar las fórmulas como reglas replicables.
    prompt = prompt_analisis(cargar_fuentes([], [ruta])[0], "Cliente")
    assert "fórmulas" in prompt and "replic" in prompt
    # Una marca de hoja filtrada dentro de un texto del RFD se detecta como las demás.
    assert MARCA_FUENTE.search("según F02 Calculo!B3 se suma")


def test_avance_se_conserva_y_se_distingue_de_las_fuentes(tmp_path):
    from docx import Document
    from generador_rfd.calidad import MARCA_FUENTE
    documento = Document()
    documento.add_paragraph("2.1 Registro de solicitudes: el usuario ingresa el pedido desde el portal.")
    ruta = tmp_path / "avance.docx"
    documento.save(ruta)

    solo = cargar_fuentes([], [], avances=[ruta])
    assert [(f.id, f.categoria) for f in solo] == [("A01", "avance")]
    con_reunion = cargar_fuentes([EJEMPLOS / "transcripcion_demo.txt"], [], avances=[ruta])
    assert [f.id for f in con_reunion] == ["F01", "A01"]

    prompt = crear_prompt(con_reunion, "", "Cliente", "C-1", "existente", "x", None)
    assert "RFD EN AVANCE" in prompt and "literalmente" in prompt
    assert prompt.count("2.1 Registro de solicitudes") == 1  # el avance no se repite en las fuentes completas
    assert "No hay transcripción de reunión" in crear_prompt(solo, "", "Cliente", "C-1", "existente", "x", None)  # un avance no es una reunión
    assert MARCA_FUENTE.search("según A01 P3")
    assert MARCA_FUENTE.search("según F02 Hoja Datos!B3 se suma")

    with pytest.raises(ValueError, match="nombres distintos"):
        cargar_fuentes([], [ruta], avances=[ruta])
    with pytest.raises(ValueError, match="al menos un archivo"):
        cargar_fuentes([], [], avances=[])


def test_excel_ignora_dimension_incorrecta_de_la_hoja(tmp_path):
    from openpyxl import Workbook
    from zipfile import ZIP_DEFLATED
    libro = Workbook()
    libro.active["A1"], libro.active["C5"] = "inicio", "fin"
    original = tmp_path / "original.xlsx"
    libro.save(original)
    roto = tmp_path / "dimension_rota.xlsx"
    with ZipFile(original) as entrada, ZipFile(roto, "w", ZIP_DEFLATED) as salida:
        for item in entrada.infolist():
            datos = entrada.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                datos = datos.replace(b'<dimension ref="A1:C5"/>', b'<dimension ref="A1"/>')
                assert b'<dimension ref="A1"/>' in datos
            salida.writestr(item, datos)
    assert cargar_fuentes([], [roto])[0].contenido.splitlines() == ["Sheet!A1: inicio", "Sheet!C5: fin"]


def _libro_vacio(ruta: Path) -> Path:
    from openpyxl import Workbook
    Workbook().save(ruta)
    return ruta


def test_transcripcion_opcional_si_hay_contexto():
    fuentes = cargar_fuentes([], [EJEMPLOS / "contexto_demo.docx", EJEMPLOS / "presentacion_demo.pptx"])
    assert [f.categoria for f in fuentes] == ["contexto", "contexto"]
    assert "No hay transcripción de reunión" in crear_prompt(fuentes, "", "Cliente", "C-1", "existente", "02/10/2026", None)
    con_reunion = crear_prompt(cargar_fuentes([EJEMPLOS / "transcripcion_demo.txt"], []), "", "C", "C-1", "nuevo", "x", None)
    assert "No hay transcripción de reunión" not in con_reunion
    with pytest.raises(ValueError, match="al menos un archivo"):
        cargar_fuentes([], [])


def test_prompt_separa_estilo_y_analisis():
    fuentes = cargar_fuentes([EJEMPLOS / "transcripcion_demo.txt"], [], [EJEMPLOS / "contexto_demo.txt"])
    prompt = crear_prompt(fuentes, "", "C", "C-1", "nuevo", "x", None, {"F01": {"resumen": "ANALISIS-UNICO"}})
    assert "ANALISIS-UNICO" in prompt
    assert prompt.index("REFERENCIAS DE ESTILO") < prompt.index("FUENTES COMPLETAS")
    assert "[FUENTE E01" not in prompt.split("FUENTES COMPLETAS")[1]
    assert temas_de_codigo({"F01": {"terminos_codigo": ["pedido", "contrato"]}}, None, "") == ["pedido contrato"]
    assert temas_de_codigo({}, None, "precios") == ["precios"]


def test_demo_v2_se_publica_sin_hallazgos_y_numera_por_capitulo(tmp_path):
    datos = json.loads((EJEMPLOS / "rfd_demo.json").read_text(encoding="utf-8"))
    datos["definiciones"] = [{"id": "D-01", "titulo": "Prueba", "descripcion": "Texto.", "evidencias": []}]
    r = publicar(datos, tmp_path / "RFD.docx")
    assert [h.codigo for h in r.hallazgos] == ["definiciones"]
    texto = texto_docx(r.salida)
    assert "Definiciones requeridas para construcción" in texto and "D-01 — Prueba" in texto
    assert "Adecuaciones de los componentes existentes" in texto
    with ZipFile(r.salida) as z:
        documento = etree.fromstring(z.read("word/document.xml"))
        numeracion = etree.fromstring(z.read("word/numbering.xml"))
    assert documento.xpath("count(//w:instrText[contains(., 'TOC')])", namespaces=NS) == 1
    # El título 2 del Capítulo 3 usa una lista cuyo nivel 0 empieza en 3 (→ «3.1.»).
    titulo = documento.xpath("//w:p[.//w:numId and .//w:t[text()='Definiciones requeridas para construcción']]", namespaces=NS)[0]
    num_id = titulo.xpath("string(.//w:numId/@w:val)", namespaces=NS)
    abstracto = numeracion.xpath(f"string(//w:num[@w:numId='{num_id}']/w:abstractNumId/@w:val)", namespaces=NS)
    inicio = numeracion.xpath(f"string(//w:abstractNum[@w:abstractNumId='{abstracto}']/w:lvl[@w:ilvl='0']/w:start/@w:val)", namespaces=NS)
    assert inicio == "3"


def test_revision_detecta_plantilla_y_contenido_heredado():
    fuentes = [Fuente("F01", "a.txt", "contexto", "Portal Contoso pedidos"),
               Fuente("E01", "ref.docx", "estilo", "Banco Northwind aprobó. Banco Northwind usa Swift.")]
    terminos = terminos_ajenos(fuentes, "Contoso")
    assert "Northwind" in terminos and "Contoso" not in terminos
    hallazgos = auditar_texto_final("Texto [Nombre del cliente] y Banco Northwind", terminos)
    assert {h.codigo for h in hallazgos} == {"instrucciones_plantilla", "referencia_heredada"}


def test_inventario_de_varios_repositorios(tmp_path):
    web = tmp_path / "repo-web" / "src" / "app"
    web.mkdir(parents=True)
    (tmp_path / "repo-web" / "angular.json").write_text("{}", encoding="utf-8")
    (web / "pedido.component.ts").write_text("", encoding="utf-8")
    db = tmp_path / "repo-db" / "sp"
    db.mkdir(parents=True)
    (tmp_path / "repo-db" / "README.md").write_text("# Base de datos", encoding="utf-8")
    (db / "p.sql").write_bytes("CREATE PROCEDURE dbo.Ped_Registrar AS SELECT 1".encode("utf-16"))
    inventario = inventario_repositorios(tmp_path)
    assert "2 repositorio(s)" in inventario
    assert "Componentes Angular (1): pedido" in inventario
    assert "dbo.Ped_Registrar" in inventario


def test_comando_codex_aislado_y_eventos(tmp_path):
    from generador_rfd.modelo import comando_codex, eventos_codex

    nube = comando_codex("codex", "codex", "gpt-6-luna", "low", tmp_path, tmp_path / "e.json", tmp_path / "r.json")
    for opcion in ("--ephemeral", "--ignore-user-config", "--ignore-rules", "--json", "--output-schema"):
        assert opcion in nube
    assert nube[nube.index("--sandbox") + 1] == "read-only"
    assert 'model_reasoning_effort="low"' in nube and "--oss" not in nube
    local = comando_codex("codex", "ollama", "llama", "high", tmp_path, tmp_path / "e.json", tmp_path / "r.json")
    assert local[local.index("--local-provider") + 1] == "ollama"
    assert not any("model_reasoning_effort" in x for x in local), "el esfuerzo sólo aplica a Codex en la nube"
    tokens, error = eventos_codex("\n".join([
        '{"type":"thread.started"}', "texto no JSON",
        '{"type":"turn.completed","usage":{"input_tokens":100,"output_tokens":7,"reasoning_output_tokens":3}}',
        '{"type":"error","message":"límite de uso"}',
    ]))
    assert tokens == {"input_tokens": 100, "output_tokens": 7, "reasoning_output_tokens": 3}
    assert error == "límite de uso"


def _rfd_de_nivel() -> dict:
    """RFD sintético que imita las proporciones del RFD de referencia aprobado."""
    frase = lambda n, base="El portal deberá registrar el resultado de la operación": " ".join([base] + ["con trazabilidad completa"] * ((n - len(base.split())) // 3 + 1)).split()[:n]  # noqa: E731
    parrafo = lambda n: " ".join(frase(n, "Se adecuará la consulta del portal para utilizar los datos vigentes")) + "."  # noqa: E731
    criterio = lambda extra="": " ".join(frase(20)) + extra + "."  # noqa: E731
    datos = json.loads((EJEMPLOS / "rfd_demo.json").read_text(encoding="utf-8"))
    datos.update(
        justificacion=[parrafo(70), parrafo(65)],
        objetivo="Adecuar " + " ".join(frase(39)),
        proceso_actual=[parrafo(40)] * 4,
        secciones_generales=[{"titulo": "Integración", "parrafos": [
            parrafo(45), "La integración deberá distinguir una consulta sin resultados, una validación rechazada y una falla de comunicación con el sistema externo."]}],
    )
    datos["alcance"]["items"] = [{"id": a["id"], "descripcion": "Registrar " + " ".join(frase(11))} for a in datos["alcance"]["items"]]
    for r in datos["requerimientos"]:
        r["parrafos"] = [parrafo(40), parrafo(35)]
        r["flujo"] = [{"numero": i, "actor": "Usuario", "componente": "Portal", "accion": "Registra la información requerida del paso"} for i in range(1, 6)]
        r["criterios_aceptacion"] = [criterio() for _ in range(4)] + [criterio(" y no deberá continuar ante un error")]
    datos["prerrequisitos"]["items"] = ["El Cliente deberá " + " ".join(frase(16)) for _ in range(9)]
    datos["restricciones"] = [" ".join(frase(16)) for _ in range(9)]
    datos["anexos"] = [{"id": "F-01", "nombre": "a.docx", "uso": "Interfaz de solicitudes; base de RF-01."}]
    return datos


def test_rubrica_aprueba_nivel_de_referencia_y_rechaza_borrador_pobre():
    from generador_rfd.calidad import evaluar, instrucciones_correccion

    bueno = evaluar(_rfd_de_nivel())
    assert bueno.aprobado(), instrucciones_correccion(bueno)
    pobre = evaluar(json.loads((EJEMPLOS / "rfd_demo.json").read_text(encoding="utf-8")))
    assert not pobre.aprobado()
    fallas = {c.id for c in pobre.fallas}
    assert {"justificacion", "voz_normativa", "RF-01.criterios"} <= fallas
    assert "[RF-01.criterios]" in instrucciones_correccion(pobre)


def test_rubrica_detecta_marcas_de_fuente_y_criterios_cortos():
    from generador_rfd.calidad import evaluar

    datos = _rfd_de_nivel()
    datos["requerimientos"][0]["criterios_aceptacion"] = ["Muestra datos.", "Funciona bien (F01 P12).", "Valida.", "Guarda."]
    fallas = {c.id for c in evaluar(datos).fallas}
    assert {"RF-01.criterios", "sin_marcas"} <= fallas


def _demo_con_preguntas() -> dict:
    datos = json.loads((EJEMPLOS / "rfd_demo.json").read_text(encoding="utf-8"))
    datos["preguntas_cliente"] = [
        {"id": "P-01", "tema": "Países", "contexto": "La especificación incluye cuatro países y otros documentos mencionan dos más.",
         "pregunta": "¿Qué países quedan incluidos en esta entrega?", "relacion": ["D-01"]},
        {"id": "P-02", "tema": "Transición", "contexto": "Hoy los pedidos se registran en Excel y luego se cargan en el sistema.",
         "pregunta": "¿El proceso actual coexistirá con la integración?", "relacion": ["RF-01"]},
    ]
    return datos


def test_cuestionario_y_seccion_de_preguntas(tmp_path):
    from docx import Document

    from generador_rfd.cuestionario import leer_cuestionario

    r = publicar(_demo_con_preguntas(), tmp_path / "RFD.docx")
    texto = texto_docx(r.salida)
    assert "Preguntas para el cliente" in texto and "1. Países: " in texto
    cuestionario = r.salida.with_suffix(".cuestionario.docx")
    assert cuestionario.is_file()
    doc = Document(cuestionario)
    tabla = doc.tables[0]
    assert [c.text for c in tabla.rows[0].cells] == ["N.º", "Tema", "Contexto y pregunta", "Respuesta del cliente"]
    tabla.rows[1].cells[3].text = "Colombia, Chile, Perú y Ecuador."
    tabla.rows[2].cells[3].text = "Pendiente"
    doc.save(cuestionario)
    assert leer_cuestionario(cuestionario) == {"P-01": "Colombia, Chile, Perú y Ecuador."}


def test_historial_sqlite_guarda_versiones_y_respuestas(tmp_path):
    from generador_rfd.almacen import Almacen

    a = Almacen(tmp_path / "rfd.db")
    r = publicar(_demo_con_preguntas(), tmp_path / "RFD.docx")
    base = a.registrar(r.datos, origen="generado", ruta_docx=r.salida, calidad={"puntaje": 0.9})
    assert a.responder(base, "P-01", "Colombia y Chile")["estado"] == "respondida"
    assert a.responder(base, "P-02", "Pendiente")["estado"] == "pendiente"
    assert a.respuestas(base) == {"P-01": "Colombia y Chile"}
    hijo = a.registrar({**r.datos, "version": "0.2", "preguntas_cliente": []}, origen="mejorado",
                       ruta_docx=tmp_path / "RFD_v0.2.docx", padre_id=base)
    filas = {f["id"]: f for f in a.listar()}
    assert filas[base]["mejorado"] and filas[base]["respondidas"] == 1 and filas[base]["calidad"] == 0.9
    assert filas[hijo]["padre_id"] == base and filas[hijo]["version"] == "0.2"
    # Un RFD anterior (sin preguntas ni versión) sigue siendo una base válida.
    viejo = json.loads((EJEMPLOS / "rfd_demo.json").read_text(encoding="utf-8"))
    (tmp_path / "salidas").mkdir()
    publicar(viejo, tmp_path / "salidas" / "RFD_VIEJO.docx")
    assert a.importar_salidas(tmp_path / "salidas") == 1


def test_mejora_sube_version_y_detecta_respuestas_sin_incorporar(tmp_path, monkeypatch):
    from generador_rfd import pipeline

    def adaptar_falso(*, datos, respuestas, adjuntos, **_):
        nuevo = json.loads(json.dumps({k: v for k, v in datos.items() if k not in ("version", "historial")}))
        nuevo["alcance"]["introduccion"] = ["El alcance comprende Colombia y Chile."]
        nuevo["preguntas_cliente"] = [p for p in nuevo["preguntas_cliente"] if p["id"] != "P-01"]
        nuevo["cambios"] = [{"pregunta": "P-01", "seccion": "1.3 Alcance", "cambio": "Se fijan los países."}]
        return nuevo

    monkeypatch.setattr(pipeline, "adaptar_rfd", adaptar_falso)
    monkeypatch.setattr(pipeline, "asegurar_calidad", lambda d, *a, **k: (d, {"puntaje": 1.0, "umbral": 0.85, "criterios": [], "rondas": 0}))
    base = pipeline.normalizar(_demo_con_preguntas())
    r = pipeline.mejorar(pipeline.OpcionesMejora(
        base=tmp_path / "RFD_DEMO.docx", salida_dir=tmp_path, respuestas={"P-01": "Colombia y Chile", "P-02": "Sí, coexistirán"},
        datos_base=base, revisar_word=False))
    assert r.salida.name == "RFD_DEMO_v0.2.docx"
    assert r.datos["version"] == "0.2" and r.datos["historial"][-1]["version"] == "0.2"
    assert [p["id"] for p in r.datos["preguntas_cliente"]] == ["P-02"]
    assert any(h.codigo == "respuesta_sin_incorporar" and "P-02" in h.mensaje for h in r.hallazgos)
    assert "0.2" in texto_docx(r.salida)


def test_reglas_del_equipo_se_inyectan_y_verifican(tmp_path, monkeypatch):
    from generador_rfd import pipeline
    from generador_rfd.almacen import Almacen
    from generador_rfd.calidad import evaluar
    from generador_rfd.modelo import prompt_adaptacion
    from generador_rfd.reglas import Regla, bloque_prompt, incumplidas

    a = Almacen(tmp_path / "rfd.db")
    evitar = a.guardar_regla(Regla(texto="No usar enumeraciones abiertas como etc.", patron="etc.|entre otros"))
    exigir = a.guardar_regla(Regla(texto="Cada RF incluye un criterio de timeout", tipo="exigir", seccion="criterios", patron="timeout"))
    otro = a.guardar_regla(Regla(texto="Sólo para otro cliente: no citar SAP", cliente="Otro SA", patron="SAP"))
    a.guardar_regla(Regla(texto="Regla desactivada de prueba", activa=False))
    with pytest.raises(ValueError, match="expresión válida"):
        a.guardar_regla(Regla(texto="Patrón roto en la verificación", patron="re:(abc"))
    reglas = a.reglas_para("Cliente Demostración")
    assert [r.id for r in reglas] == [evitar.id, exigir.id]
    assert otro.id in [r.id for r in a.reglas_para("otro sa")]

    datos = pipeline.normalizar(_demo_con_preguntas())
    rf = datos["requerimientos"][0]
    rf["criterios_aceptacion"][0] += " Aplica a solicitudes, reclamos, etc."
    fallas = dict((r.id, d) for r, d in incumplidas(datos, reglas))
    assert "etc." in fallas[evitar.id] and rf["id"] in fallas[evitar.id]
    assert fallas[exigir.id].startswith("falta en") and rf["id"] in fallas[exigir.id]
    ev = {c.id: c for c in evaluar(datos, None, False, reglas).criterios}
    assert not ev[f"regla.R{evitar.id}"].ok and not ev[f"regla.R{exigir.id}"].ok

    bloque = bloque_prompt(reglas)
    assert f"[R{evitar.id}] NO HACER (Todo el documento)" in bloque and f"[R{exigir.id}] HACER SIEMPRE (Criterios" in bloque
    prompt = crear_prompt(cargar_fuentes([EJEMPLOS / "transcripcion_demo.txt"], []), "", "C", "C-1", "nuevo", "x", None, None, bloque)
    assert prompt.index("REGLAS DE INTEGRIDAD") < prompt.index("REGLAS DEL EQUIPO") < prompt.index("CÓMO REDACTAR")
    assert "REGLAS DEL EQUIPO" in prompt_adaptacion(datos, [], [], bloque)

    recibido = {}

    def adaptar_falso(*, datos, reglas="", **_):
        recibido["reglas"] = reglas
        nuevo = json.loads(json.dumps({k: v for k, v in datos.items() if k not in ("version", "historial")}))
        nuevo["cambios"] = [{"pregunta": "P-01", "seccion": "1.3 Alcance", "cambio": "Se fijan los países."}]
        return nuevo

    monkeypatch.setattr(pipeline, "adaptar_rfd", adaptar_falso)
    monkeypatch.setattr(pipeline, "asegurar_calidad", lambda d, *a, **k: (d, {"puntaje": 1.0, "umbral": 0.85, "criterios": [], "rondas": 0}))
    r = pipeline.mejorar(pipeline.OpcionesMejora(base=tmp_path / "RFD_DEMO.docx", salida_dir=tmp_path, respuestas={"P-01": "Chile"},
                                                 datos_base=datos, revisar_word=False, reglas=reglas))
    assert "REGLAS DEL EQUIPO" in recibido["reglas"]
    assert {h.codigo for h in r.hallazgos} >= {"regla_incumplida"}


requiere_codex = pytest.mark.skipif(not shutil.which("codex"), reason="Codex CLI no está en PATH")


@pytest.mark.codex
@requiere_codex
def test_codex_analiza_un_documento(tmp_path):
    from generador_rfd.modelo import Uso

    uso = Uso()
    fuentes = cargar_fuentes([], [EJEMPLOS / "contexto_demo.docx"])
    analisis = analizar_fuentes(fuentes, cliente="Cliente prueba", proveedor="codex",
                                modelo=os.environ.get("RFD_MODELO_CODEX", "gpt-6-luna"), esfuerzo="low",
                                directorio=tmp_path, uso=uso)
    assert analisis["F01"]["resumen"]
    assert uso.llamadas == 1 and uso.tokens_entrada > 0
    assert not list(tmp_path.iterdir()), "el directorio temporal de Codex debe eliminarse"


@pytest.mark.claude
@requiere_claude
def test_cli_devuelve_salida_estructurada(tmp_path):
    esquema = {
        "type": "object",
        "properties": {"estado": {"type": "string", "enum": ["ok"]}, "suma": {"type": "integer"}},
        "required": ["estado", "suma"],
        "additionalProperties": False,
    }
    p = subprocess.run(
        comando_claude(CLAUDE, MODELO, esquema),
        input="Responde estado=ok y suma=2+3.", text=True, encoding="utf-8",
        capture_output=True, timeout=180, cwd=tmp_path, check=False,
    )
    assert p.returncode == 0, p.stderr[-800:]
    respuesta = json.loads(p.stdout[p.stdout.find("{"):])
    assert respuesta["is_error"] is False, respuesta.get("result")
    assert respuesta["structured_output"] == {"estado": "ok", "suma": 5}
    uso = respuesta["usage"]
    assert uso["input_tokens"] + uso["cache_creation_input_tokens"] + uso["cache_read_input_tokens"] < 30_000


@pytest.mark.claude
@requiere_claude
def test_analisis_por_documento(tmp_path):
    fuentes = cargar_fuentes([], [EJEMPLOS / "contexto_demo.docx"])
    analisis = analizar_fuentes(fuentes, cliente="Cliente prueba", proveedor="claude", modelo=MODELO, directorio=tmp_path)
    a = analisis["F01"]
    assert a["resumen"] and a["terminos_codigo"]
    assert all(e.startswith("F01") for e in a["evidencias"])
    assert not list(tmp_path.iterdir()), "el directorio temporal de Claude debe eliminarse"


@pytest.mark.claude
@requiere_claude
def test_api_genera_word_con_varios_archivos(tmp_path, monkeypatch):
    from fastapi.testclient import TestClient

    monkeypatch.setattr(servidor, "SALIDAS", tmp_path)
    cliente = TestClient(servidor.app)
    archivos = [
        ("transcripciones", ("transcripcion_demo.txt", (EJEMPLOS / "transcripcion_demo.txt").read_bytes())),
        ("contextos", ("contexto_demo.docx", (EJEMPLOS / "contexto_demo.docx").read_bytes())),
        ("contextos", ("presentacion_demo.pptx", (EJEMPLOS / "presentacion_demo.pptx").read_bytes())),
    ]
    r = cliente.post("/api/generar", files=archivos, data={
        "cliente": "Cliente prueba", "codigo": "TEST-API", "tipo_proyecto": "existente",
        "proveedor": "claude", "modelo": MODELO, "revisar_word": "false",
    })
    assert r.status_code == 200, r.text
    id_ = r.json()["id"]
    limite = time.time() + 1200
    while (t := cliente.get(f"/api/trabajos/{id_}").json())["estado"] == "en_curso":
        assert time.time() < limite, "la generación superó 20 minutos"
        time.sleep(3)
    assert t["estado"] == "listo", t.get("error")
    assert [f["nombre"] for f in t["fuentes"]] == ["transcripcion_demo.txt", "contexto_demo.docx", "presentacion_demo.pptx"]
    assert t["uso"]["llamadas"] >= 4, "tres análisis, una redacción y el detalle de cada opción"
    word = cliente.get(f"/api/trabajos/{id_}/archivo/docx")
    assert word.status_code == 200 and word.content[:2] == b"PK"
    assert cliente.get(f"/api/trabajos/{id_}/archivo/analisis").status_code == 200
    assert (tmp_path / "RFD_TEST-API.revision.md").is_file()


# ---------------------------------------------------------------- Importar un RFD en avance y continuarlo

def _docx_con_texto(ruta: Path, *parrafos: str) -> Path:
    from docx import Document

    ruta.parent.mkdir(parents=True, exist_ok=True)
    documento = Document()
    for p in parrafos:
        documento.add_paragraph(p)
    documento.save(ruta)
    return ruta


def _demo_normalizado() -> dict:
    from generador_rfd.schema import normalizar

    return normalizar(_demo_con_preguntas())


def _sin_version(datos: dict) -> dict:
    return json.loads(json.dumps({k: v for k, v in datos.items() if k not in ("version", "historial")}))


def _calidad_falsa(puntaje: float = 1.0, criterios: list[dict] | None = None, vista: dict | None = None):
    def falsa(datos, o, analisis, uso, directorio, avisar, con_codigo=None, preservar=""):
        if vista is not None:
            vista["preservar"] = preservar
        return datos, {"puntaje": puntaje, "umbral": 0.85, "rondas": 0, "criterios": criterios or []}
    return falsa


def test_correccion_recibe_el_texto_del_autor_y_prohibe_reescribirlo():
    from generador_rfd.modelo import prompt_correccion

    datos = _demo_normalizado()
    fallas = "- [RF-01.criterios] Criterios cortos. Situación actual: 3 palabras"
    assert "TEXTO DEL AUTOR" not in prompt_correccion(datos, fallas, None)
    prompt = prompt_correccion(datos, fallas, None, preservar="El portal valida el pedido. TEXTO-UNICO-DEL-AUTOR")
    assert "TEXTO-UNICO-DEL-AUTOR" in prompt
    assert "TEXTO DEL AUTOR (NO MODIFICAR)" in prompt and "agregó o completó" in prompt and "hallazgo" in prompt
    assert prompt.index("TEXTO DEL AUTOR") < prompt.index("RFD ACTUAL")


def test_asegurar_calidad_entrega_el_texto_del_autor_al_corrector(tmp_path, monkeypatch):
    from types import SimpleNamespace

    from generador_rfd import pipeline
    from generador_rfd.modelo import Uso

    datos = _demo_normalizado()
    datos["requerimientos"][0]["criterios_aceptacion"] = ["Muestra datos.", "Valida.", "Guarda.", "Envía."]
    recibido: list[str] = []

    def corregir_falso(*, datos, preservar="", **_):
        recibido.append(preservar)
        return json.loads(json.dumps(datos))

    monkeypatch.setattr(pipeline, "corregir_rfd", corregir_falso)
    o = SimpleNamespace(proveedor="claude", modelo="", esfuerzo="", rondas_calidad=1, reglas=[], proyecto=None)
    pipeline.asegurar_calidad(datos, o, None, Uso(), tmp_path, lambda *a: None, preservar="TEXTO-AUTOR")
    pipeline.asegurar_calidad(datos, o, None, Uso(), tmp_path, lambda *a: None)
    assert recibido == ["TEXTO-AUTOR", ""]


def test_generar_con_avance_protege_el_texto_del_autor_en_la_correccion(tmp_path, monkeypatch):
    from generador_rfd import pipeline

    avance = _docx_con_texto(tmp_path / "avance.docx", "2.1 Registro de solicitudes: TEXTO-AUTOR-UNICO")
    monkeypatch.setattr(pipeline, "analizar_fuentes", lambda *a, **k: {})
    monkeypatch.setattr(pipeline, "redactar_rfd", lambda **k: _demo_normalizado())
    vista: dict = {}
    monkeypatch.setattr(pipeline, "asegurar_calidad", _calidad_falsa(vista=vista))

    def opciones(**extra):
        return pipeline.Opciones(transcripciones=[], contextos=[EJEMPLOS / "contexto_demo.docx"], cliente="Cliente", codigo="C-1",
                                 tipo_proyecto="existente", fecha="01/01/2026", salida=tmp_path / "RFD.docx", revisar_word=False, **extra)

    pipeline.generar(opciones(avances=[avance]))
    assert "TEXTO-AUTOR-UNICO" in vista["preservar"]
    pipeline.generar(opciones())
    assert vista["preservar"] == ""


def test_prompt_continuacion_conserva_lo_escrito_y_no_exige_respuestas():
    from generador_rfd.modelo import prompt_continuacion

    datos = _demo_normalizado()
    contexto = [Fuente("C01", "acta.txt", "contexto", "L1: CONTEXTO-NUEVO")]
    prompt = prompt_continuacion(datos, contexto, "Enfócate en los criterios del RF-02")
    for frase in ("literalmente", "«pendientes»", "«definiciones»", "por confirmar", "RF-xx", "terminología"):
        assert frase in prompt
    assert "CONTEXTO-NUEVO" in prompt and "Enfócate en los criterios del RF-02" in prompt
    assert "RESPUESTAS DEL CLIENTE" not in prompt
    from generador_rfd.calidad import MARCA_FUENTE
    assert MARCA_FUENTE.search("según C01 P3")  # las marcas del contexto adicional tampoco van en los textos
    vacio = prompt_continuacion(datos, [], "")
    assert "(sin documentos adicionales)" in vacio and "(sin instrucción)" in vacio


def test_continuar_completa_sin_respuestas_y_reporta_lo_que_no_cumple(tmp_path, monkeypatch):
    from generador_rfd import pipeline

    base = _demo_normalizado()
    contexto = tmp_path / "contexto.txt"
    contexto.write_text("El portal también debe validar el país.", encoding="utf-8")
    recibido: dict = {}

    def continuar_falso(*, datos, contexto, instruccion, **_):
        recibido.update(contexto=contexto, instruccion=instruccion)
        nuevo = _sin_version(datos)
        nuevo["alcance"]["introduccion"] = ["El alcance se completó con el contexto adicional."]
        nuevo["cambios"] = [{"pregunta": "Continuación", "seccion": "1.3 Alcance", "cambio": "Se completa el alcance."}]
        return nuevo

    criterios = [{"id": "RF-01.criterios", "descripcion": "Criterios de 15 a 45 palabras", "ok": False, "detalle": "3 palabras"},
                 {"id": "sin_marcas", "descripcion": "Sin marcas de fuente", "ok": True, "detalle": ""}]
    monkeypatch.setattr(pipeline, "continuar_rfd", continuar_falso)
    monkeypatch.setattr(pipeline, "asegurar_calidad", _calidad_falsa(0.5, criterios, recibido))
    r = pipeline.continuar(pipeline.OpcionesMejora(base=tmp_path / "RFD_DEMO.docx", salida_dir=tmp_path, adjuntos=[contexto],
                                                   instruccion="Enfócate en RF-02", datos_base=base, revisar_word=False))
    assert r.salida.name == "RFD_DEMO_v0.2.docx"
    assert r.datos["version"] == "0.2" and "ontinuación" in r.datos["historial"][-1]["descripcion"]
    assert [f.nombre for f in recibido["contexto"]] == ["contexto.txt"] and recibido["instruccion"] == "Enfócate en RF-02"
    assert json.dumps(base["titulo"], ensure_ascii=False)[1:-1] in recibido["preservar"]
    assert any(h.codigo == "calidad_texto_autor" and "RF-01.criterios" in h.mensaje for h in r.hallazgos)
    assert not any("sin_marcas" in h.mensaje for h in r.hallazgos)
    # Sin contexto ni instrucción también es válido: no depende de respuestas del cliente.
    sola = pipeline.continuar(pipeline.OpcionesMejora(base=tmp_path / "RFD_DEMO.docx", salida_dir=tmp_path, datos_base=base, revisar_word=False))
    assert sola.datos["version"] == "0.2" and recibido["contexto"] == []


def test_continuar_avisa_si_la_ia_elimina_un_requerimiento_del_autor(tmp_path, monkeypatch):
    from generador_rfd import pipeline

    base = _demo_normalizado()

    def continuar_falso(*, datos, **_):
        nuevo = _sin_version(datos)
        nuevo["requerimientos"] = nuevo["requerimientos"][1:]
        return nuevo

    monkeypatch.setattr(pipeline, "continuar_rfd", continuar_falso)
    monkeypatch.setattr(pipeline, "asegurar_calidad", _calidad_falsa())
    r = pipeline.continuar(pipeline.OpcionesMejora(base=tmp_path / "RFD_DEMO.docx", salida_dir=tmp_path, datos_base=base, revisar_word=False))
    assert any(h.codigo == "rf_eliminado" and base["requerimientos"][0]["id"] in h.mensaje for h in r.hallazgos)


# ---------------------------------------------------------------- Especificación detallada por opción

def _vista_portal_publico() -> dict:
    return {
        "portal": "Portal Público",
        "definiciones_generales": ["Permitirá al proveedor consultar los anticipos asociados a sus Órdenes de Compra."],
        "ventanas": [{
            "nombre": "principal",
            "descripcion": "La ventana principal estará compuesta por el panel de filtros y la matriz de resultados.",
            "secciones": [{"nombre": "Matriz de resultados", "descripcion": "Presenta un registro por anticipo aplicado.",
                           "campos": [{"nombre": "Saldo de anticipo", "tipo": "Moneda (calculado)",
                                       "descripcion": "Valor del anticipo menos el valor aplicado, acumulado por Orden de Compra."}]}],
        }],
        "reglas_negocio": ["Los anticipos no podrán cruzarse entre empresas diferentes."],
    }


def _con_especificacion() -> dict:
    datos = _demo_normalizado()
    datos["requerimientos"][0]["vistas"] = [_vista_portal_publico()]
    datos["requerimientos"][0]["integraciones"] = [
        {"servicio": "Servicio de consulta de anticipos", "descripcion": "Retorna los anticipos por Orden de Compra."}]
    return datos


def test_esquema_admite_especificacion_y_normaliza_un_rfd_anterior():
    from generador_rfd.schema import normalizar, validar_estructura

    anterior = json.loads((EJEMPLOS / "rfd_demo.json").read_text(encoding="utf-8"))
    for r in anterior["requerimientos"]:
        r.pop("vistas", None), r.pop("integraciones", None)
    datos = normalizar(anterior)
    validar_estructura(datos)
    assert all(r["vistas"] == [] and r["integraciones"] == [] for r in datos["requerimientos"])
    validar_estructura(_con_especificacion())


def test_word_presenta_la_especificacion_detallada_de_la_opcion(tmp_path):
    r = publicar(_con_especificacion(), tmp_path / "RFD.docx")
    texto = texto_docx(r.salida)
    for frase in ("Definiciones generales (Portal Público)", "Definiciones de la ventana principal (Portal Público)",
                  "Matriz de resultados", "Saldo de anticipo", "Moneda (calculado)", "Reglas de negocio (Portal Público)",
                  "Los anticipos no podrán cruzarse entre empresas diferentes.", "Integraciones", "Servicio de consulta de anticipos"):
        assert frase in texto, frase
    assert texto.index("Reglas de negocio (Portal Público)") < texto.index("Flujo propuesto")
    assert not any(h.codigo == "especificacion_incompleta" for h in r.hallazgos)


def test_especificacion_sin_reglas_ni_ventanas_se_informa():
    from generador_rfd.schema import auditar_contenido

    datos = _con_especificacion()
    datos["requerimientos"][0]["vistas"][0]["reglas_negocio"] = []
    assert any(h.codigo == "especificacion_incompleta" and "RF-01" in h.mensaje for h in auditar_contenido(datos))


def test_prompt_de_detalle_se_enfoca_en_una_opcion_e_imita_al_autor():
    from generador_rfd.modelo import prompt_detalle

    datos = _demo_normalizado()
    datos["secciones_generales"] = [{"titulo": "2.1.1. Consulta de Facturas", "parrafos": ["MODELO-DEL-AUTOR con ventana principal."]}]
    fuentes = [Fuente("C01", "transcripcion.docx", "contexto", "L1: TRANSCRIPCION-REUNION")]
    prompt = prompt_detalle(datos, datos["requerimientos"][1], fuentes, instruccion="Detalla Control de Anticipos")
    assert datos["requerimientos"][1]["nombre"] in prompt
    for texto in ("MODELO-DEL-AUTOR", "TRANSCRIPCION-REUNION", "Detalla Control de Anticipos", "secciones_reemplazadas",
                  "reglas_negocio", datos["requerimientos"][0]["nombre"]):
        assert texto in prompt, texto
    assert prompt.index("OPCIÓN A DETALLAR") < prompt.index("FUENTES COMPLETAS")


def _base_con_resumen(extension: int) -> dict:
    base = _demo_normalizado()
    base["secciones_generales"].append({"titulo": "2.1.4. Opción Registro de solicitudes", "parrafos": [" ".join(["resumen"] * extension)]})
    base["requerimientos"][3]["vistas"] = [_vista_portal_publico()]
    return base


def _continuar_con_detalle(tmp_path, monkeypatch, base: dict, detalle):
    from generador_rfd import pipeline

    llamados: list[str] = []

    def detallar_falso(*, requerimiento, fuentes, **_):
        llamados.append(requerimiento["id"])
        return detalle(requerimiento)

    contexto = tmp_path / "transcripcion.txt"
    contexto.write_text("Se definió el saldo de anticipo por Orden de Compra.", encoding="utf-8")
    monkeypatch.setattr(pipeline, "continuar_rfd", lambda *, datos, **_: _sin_version(datos))
    monkeypatch.setattr(pipeline, "asegurar_calidad", _calidad_falsa())
    monkeypatch.setattr(pipeline, "detallar_requerimiento", detallar_falso)
    r = pipeline.continuar(pipeline.OpcionesMejora(base=tmp_path / "RFD_DEMO.docx", salida_dir=tmp_path, adjuntos=[contexto],
                                                   datos_base=base, revisar_word=False))
    return r, llamados


def test_continuar_detalla_cada_opcion_y_reemplaza_su_resumen(tmp_path, monkeypatch):
    def detalle(req):
        if req["id"] != "RF-01":
            return {"vistas": [], "integraciones": [], "secciones_reemplazadas": []}
        return {"vistas": [_vista_portal_publico()], "integraciones": [],
                "secciones_reemplazadas": ["2.1.4. Opción Registro de solicitudes", "Título inexistente"]}

    r, llamados = _continuar_con_detalle(tmp_path, monkeypatch, _base_con_resumen(8), detalle)
    assert llamados == ["RF-01", "RF-02", "RF-03"], "una llamada por opción sin especificación"
    assert r.datos["requerimientos"][0]["vistas"][0]["portal"] == "Portal Público"
    assert "2.1.4. Opción Registro de solicitudes" not in [s["titulo"] for s in r.datos["secciones_generales"]]
    assert any(c["seccion"].startswith("RF-01") and "especificación detallada" in c["cambio"] for c in r.datos["cambios"])
    assert "Saldo de anticipo" in texto_docx(r.salida)


def test_el_detalle_no_reemplaza_una_seccion_mas_extensa_que_el(tmp_path, monkeypatch):
    def detalle(req):
        return {"vistas": [_vista_portal_publico()], "integraciones": [],
                "secciones_reemplazadas": ["2.1.4. Opción Registro de solicitudes"]}

    r, _ = _continuar_con_detalle(tmp_path, monkeypatch, _base_con_resumen(400), detalle)
    assert "2.1.4. Opción Registro de solicitudes" in [s["titulo"] for s in r.datos["secciones_generales"]]


def test_una_falla_al_detallar_no_bloquea_la_version(tmp_path, monkeypatch):
    def detalle(req):
        raise RuntimeError("tiempo agotado")

    r, llamados = _continuar_con_detalle(tmp_path, monkeypatch, _base_con_resumen(8), detalle)
    assert r.salida.is_file() and len(llamados) == 3
    assert sum(h.codigo == "detalle_no_generado" for h in r.hallazgos) == 3


def test_la_revision_editorial_no_altera_la_especificacion(tmp_path, monkeypatch):
    from types import SimpleNamespace

    from generador_rfd import pipeline
    from generador_rfd.modelo import Uso

    datos = _con_especificacion()
    datos["requerimientos"][0]["criterios_aceptacion"] = ["Muestra datos.", "Valida.", "Guarda.", "Envía."]

    def corregir_falso(*, datos, **_):
        nuevo = json.loads(json.dumps(datos))
        nuevo["requerimientos"][0]["vistas"] = []
        nuevo["requerimientos"][0]["integraciones"] = []
        return nuevo

    monkeypatch.setattr(pipeline, "corregir_rfd", corregir_falso)
    o = SimpleNamespace(proveedor="claude", modelo="", esfuerzo="", rondas_calidad=1, reglas=[], proyecto=None)
    mejor, _ = pipeline.asegurar_calidad(datos, o, None, Uso(), tmp_path, lambda *a: None)
    assert mejor["requerimientos"][0]["vistas"] == [_vista_portal_publico()]
    assert mejor["requerimientos"][0]["integraciones"][0]["servicio"] == "Servicio de consulta de anticipos"


def test_importar_rechaza_archivos_invalidos(tmp_path):
    from docx import Document

    from generador_rfd import pipeline

    texto = tmp_path / "avance.txt"
    texto.write_text("hola", encoding="utf-8")
    roto = tmp_path / "roto.docx"
    roto.write_bytes(b"esto no es un zip")
    vacio = tmp_path / "vacio.docx"
    Document().save(vacio)
    with pytest.raises(ValueError, match=r"Sólo se admite.*\.docx"):
        pipeline.leer_avance(texto)
    with pytest.raises(ValueError, match="No se pudo leer"):
        pipeline.leer_avance(roto)
    with pytest.raises(ValueError, match="no contiene texto"):
        pipeline.leer_avance(vacio)
    assert "Capítulo 1" in pipeline.leer_avance(_docx_con_texto(tmp_path / "bien.docx", "Capítulo 1"))


def test_importar_reconstruye_el_word_y_lo_deja_enlazado(tmp_path, monkeypatch):
    from generador_rfd import pipeline

    original = _docx_con_texto(tmp_path / "entrada" / "mi avance.docx", "(RFD-9)", "Portal de pedidos", "2.1 Registro de solicitudes")
    visto: dict = {}

    def reconstruir_falso(*, texto, cliente, codigo, **_):
        visto.update(texto=texto, cliente=cliente, codigo=codigo)
        return _demo_normalizado()

    monkeypatch.setattr(pipeline, "reconstruir_rfd", reconstruir_falso)

    def importar():
        return pipeline.importar(pipeline.OpcionesImportacion(archivo=original, salida_dir=tmp_path / "salidas", cliente="Cliente X",
                                                              codigo="RFD-9", titulo="Mi título"))

    r = importar()
    assert visto["cliente"] == "Cliente X" and visto["codigo"] == "RFD-9" and "2.1 Registro" in visto["texto"]
    assert r.salida.name == "RFD_RFD-9_importado.docx" and r.salida.read_bytes() == original.read_bytes()
    assert r.datos["titulo"] == "Mi título" and r.datos["version"] == "0.1"
    assert r.salida.with_suffix(".json").is_file() and r.calidad["criterios"]
    # Importar de nuevo el mismo código no pisa el registro anterior.
    segunda = importar()
    assert segunda.salida != r.salida and r.salida.is_file() and segunda.salida.is_file()


@pytest.fixture
def api(tmp_path, monkeypatch):
    from fastapi.testclient import TestClient

    from generador_rfd.almacen import Almacen

    monkeypatch.setattr(servidor, "SALIDAS", tmp_path / "salidas")
    monkeypatch.setattr(servidor, "_almacen", Almacen(tmp_path / "rfd.db"))
    return TestClient(servidor.app)


def _esperar(cliente, id_: str) -> dict:
    limite = time.time() + 30
    while (t := cliente.get(f"/api/trabajos/{id_}").json())["estado"] == "en_curso":
        assert time.time() < limite, "el trabajo no terminó"
        time.sleep(0.05)
    return t


def _importar_por_api(cliente, tmp_path, monkeypatch) -> tuple[int, bytes]:
    from generador_rfd import pipeline

    contenido = _docx_con_texto(tmp_path / "avance.docx", "(RFD-9)", "2.1 Registro de solicitudes").read_bytes()
    monkeypatch.setattr(pipeline, "reconstruir_rfd", lambda *, cliente, codigo, **k: {**_demo_normalizado(), "cliente": cliente, "codigo": codigo})
    r = cliente.post("/api/rfds/importar", files=[("archivo", ("avance.docx", contenido))], data={
        "cliente": "Cliente X", "codigo": "RFD-9", "proveedor": "claude", "modelo": "haiku"})
    assert r.status_code == 200, r.text
    t = _esperar(cliente, r.json()["id"])
    assert t["estado"] == "listo", t.get("error")
    return t["resultado"]["rfd_id"], contenido


def test_api_importa_un_avance_al_historial(api, tmp_path, monkeypatch):
    rfd_id, contenido = _importar_por_api(api, tmp_path, monkeypatch)
    fila = next(f for f in api.get("/api/rfds").json() if f["id"] == rfd_id)
    assert fila["origen"] == "en_avance" and fila["codigo"] == "RFD-9" and fila["padre_id"] is None
    detalle = api.get(f"/api/rfds/{rfd_id}").json()
    assert detalle["archivos"]["docx"] is True and detalle["origen"] == "en_avance"
    assert api.get(f"/api/rfds/{rfd_id}/archivo/docx").content == contenido


def test_api_importar_valida_la_entrada(api, tmp_path):
    from docx import Document

    valido = _docx_con_texto(tmp_path / "ok.docx", "Capítulo 1").read_bytes()
    vacio = tmp_path / "vacio.docx"
    Document().save(vacio)
    datos = {"cliente": "Cliente X", "codigo": "RFD-9"}

    def importar(nombre, contenido, **cambios):
        return api.post("/api/rfds/importar", files=[("archivo", (nombre, contenido))], data={**datos, **cambios})

    assert "Sólo se admite" in importar("avance.pdf", b"%PDF").json()["detail"]
    assert "No se pudo leer" in importar("roto.docx", b"no es zip").json()["detail"]
    assert "no contiene texto" in importar("vacio.docx", vacio.read_bytes()).json()["detail"]
    assert importar("ok.docx", valido, codigo="  ").status_code == 400
    assert importar("ok.docx", valido, cliente=" ").status_code == 400
    assert api.get("/api/rfds").json() == []


def test_api_continua_un_rfd_importado_sin_respuestas(api, tmp_path, monkeypatch):
    from generador_rfd import pipeline

    rfd_id, _ = _importar_por_api(api, tmp_path, monkeypatch)
    recibido: dict = {}

    def continuar_falso(*, datos, contexto, instruccion, **_):
        recibido.update(contexto=[f.nombre for f in contexto], instruccion=instruccion)
        nuevo = _sin_version(datos)
        nuevo["cambios"] = [{"pregunta": "Continuación", "seccion": "1.3 Alcance", "cambio": "Se completa."}]
        return nuevo

    monkeypatch.setattr(pipeline, "continuar_rfd", continuar_falso)
    monkeypatch.setattr(pipeline, "asegurar_calidad", _calidad_falsa())
    assert api.post("/api/rfds/999/continuar", data={"proveedor": "claude"}).status_code == 404
    r = api.post(f"/api/rfds/{rfd_id}/continuar", files=[("adjuntos", ("acta.txt", b"Acta de la reunion"))],
                 data={"instruccion": "Enfócate en RF-02", "proveedor": "claude", "modelo": "haiku", "revisar_word": "false"})
    assert r.status_code == 200, r.text
    t = _esperar(api, r.json()["id"])
    assert t["estado"] == "listo", t.get("error")
    assert recibido == {"contexto": ["acta.txt"], "instruccion": "Enfócate en RF-02"}
    nuevo = api.get(f"/api/rfds/{t['resultado']['rfd_id']}").json()
    assert nuevo["padre_id"] == rfd_id and nuevo["origen"] == "continuado" and nuevo["version"] == "0.2"
    assert api.get(f"/api/rfds/{rfd_id}").json()["version"] == "0.1"
    assert next(f for f in api.get("/api/rfds").json() if f["id"] == rfd_id)["mejorado"] is True


def test_continuar_un_importado_no_pisa_al_rfd_generado_ni_hereda_su_analisis(api, tmp_path, monkeypatch):
    from generador_rfd import pipeline

    rfd_id, _ = _importar_por_api(api, tmp_path, monkeypatch)
    ajeno = servidor.SALIDAS / "RFD_RFD-9.analisis.json"  # análisis de otro RFD con el mismo código
    ajeno.write_text(json.dumps({"F01": {"resumen": "de otro RFD"}}), encoding="utf-8")
    recibido: dict = {}
    calidad_falsa = _calidad_falsa()

    def calidad_espia(nuevo, o, analisis, *a, **k):
        recibido["analisis"] = analisis
        return calidad_falsa(nuevo, o, analisis, *a, **k)

    monkeypatch.setattr(pipeline, "continuar_rfd", lambda *, datos, **_: _sin_version(datos))
    monkeypatch.setattr(pipeline, "asegurar_calidad", calidad_espia)
    r = api.post(f"/api/rfds/{rfd_id}/continuar", data={"proveedor": "claude", "modelo": "haiku", "revisar_word": "false"})
    t = _esperar(api, r.json()["id"])
    assert t["estado"] == "listo", t.get("error")
    nuevo = servidor.almacen().obtener(t["resultado"]["rfd_id"])
    assert Path(nuevo["ruta_docx"]).name == "RFD_RFD-9_importado_v0.2.docx"
    assert recibido["analisis"] is None
    assert not (servidor.SALIDAS / "RFD_RFD-9_importado_v0.2.analisis.json").exists()
    assert ajeno.is_file()


def test_api_importar_con_nombre_peligroso_no_deja_trabajos_huerfanos(api):
    antes = dict(servidor._trabajos)
    r = api.post("/api/rfds/importar", files=[("archivo", ("..", b"x"))], data={"cliente": "Cliente X", "codigo": "RFD-9"})
    assert r.status_code == 400
    assert servidor._trabajos == antes


def test_cli_importa_y_continua(api, tmp_path, monkeypatch):
    from generador_rfd import cli, pipeline

    a = cli.parser().parse_args(["importar", "avance.docx", "--cliente", "C", "--codigo", "K"])
    assert a.comando == "importar" and a.archivo == Path("avance.docx") and a.titulo is None
    b = cli.parser().parse_args(["continuar", "--id", "3", "--contexto", "a.txt", "--instruccion", "foco"])
    assert (b.comando, b.id, b.contexto, b.instruccion) == ("continuar", 3, [Path("a.txt")], "foco")

    original = _docx_con_texto(tmp_path / "avance.docx", "(K-1)", "2.1 Registro")
    monkeypatch.setattr(pipeline, "reconstruir_rfd", lambda **k: _demo_normalizado())
    monkeypatch.setattr(pipeline, "continuar_rfd", lambda *, datos, **_: _sin_version(datos))
    monkeypatch.setattr(pipeline, "asegurar_calidad", _calidad_falsa())
    assert cli.main(["importar", str(original), "--cliente", "Cliente X", "--codigo", "K-1"]) == 0
    assert cli.main(["importar", str(tmp_path / "no_existe.docx"), "--cliente", "C", "--codigo", "K"]) == 2
    (importado,) = servidor.almacen().listar()
    assert importado["origen"] == "en_avance"
    assert cli.main(["continuar", "--id", str(importado["id"]), "--sin-word"]) == 0
    versiones = {f["version"]: f for f in servidor.almacen().listar()}
    assert versiones["0.2"]["origen"] == "continuado" and versiones["0.2"]["padre_id"] == importado["id"]


# ---------------------------------------------------------------- Prototipo desde el proyecto original

def _proyecto_web(tmp_path, **extra) -> Path:
    raiz = tmp_path / "proyecto"
    raiz.mkdir()
    (raiz / "package.json").write_text(json.dumps({"scripts": {"dev": "vite", "build": "vite build", "start": "node s.js"}}), encoding="utf-8")
    for nombre, contenido in extra.items():
        destino = raiz / nombre
        destino.parent.mkdir(parents=True, exist_ok=True)
        destino.write_text(contenido, encoding="utf-8")
    return raiz


def test_detecta_comandos_sin_ejecutar_nada(tmp_path, monkeypatch):
    from generador_rfd import prototipo

    monkeypatch.setattr(prototipo.shutil, "which", lambda nombre: None)
    ejecutado = []
    monkeypatch.setattr(prototipo.subprocess, "Popen", lambda *a, **k: ejecutado.append(a))
    monkeypatch.setattr(prototipo.subprocess, "run", lambda *a, **k: ejecutado.append(a))
    raiz = _proyecto_web(tmp_path, **{"bun.lock": "", "Makefile": "dev:\n\tbun run dev\nbuild:\n\techo\nrun: dev\n",
                                       "manage.py": "", "index.html": "<html></html>"})
    comandos = prototipo.detectar_comandos(raiz)
    por_etiqueta = {c["etiqueta"]: c for c in comandos}
    assert por_etiqueta["bun run dev"]["comando"] == ["bun", "run", "dev"]
    assert por_etiqueta["bun run dev"]["url_esperada"] == "http://localhost:5173"
    assert por_etiqueta["bun run dev"]["cwd"] == str(raiz)
    assert "bun run start" in por_etiqueta and "bun run build" not in por_etiqueta
    assert por_etiqueta["make dev"]["comando"] == ["make", "dev"] and "make run" in por_etiqueta
    assert por_etiqueta["python manage.py runserver"]["url_esperada"] == "http://127.0.0.1:8000"
    estatico = next(c for c in comandos if "http.server" in c["etiqueta"])
    assert estatico["comando"][1:4] == ["-m", "http.server", estatico["comando"][3]] and estatico["url_esperada"].startswith("http://127.0.0.1:")
    assert ejecutado == []


@pytest.mark.parametrize("candado,esperado", [("pnpm-lock.yaml", ["pnpm", "run"]), ("yarn.lock", ["yarn"]), ("", ["npm", "run"])])
def test_el_gestor_de_paquetes_sale_del_lockfile(tmp_path, monkeypatch, candado, esperado):
    from generador_rfd import prototipo

    monkeypatch.setattr(prototipo.shutil, "which", lambda nombre: None)
    raiz = _proyecto_web(tmp_path, **({candado: ""} if candado else {}))
    assert next(c for c in prototipo.detectar_comandos(raiz) if c["etiqueta"].endswith(" dev"))["comando"][:len(esperado)] == esperado


def test_resuelve_el_shim_cmd_de_windows(tmp_path, monkeypatch):
    from generador_rfd import prototipo

    monkeypatch.setattr(prototipo.shutil, "which", lambda nombre: r"C:\bun\bun.CMD" if nombre == "bun" else None)
    raiz = _proyecto_web(tmp_path, **{"bun.lock": ""})
    c = next(c for c in prototipo.detectar_comandos(raiz) if c["etiqueta"] == "bun run dev")
    assert c["comando"][0] == r"C:\bun\bun.CMD"


def test_detecta_en_subcarpetas_y_exige_directorio(tmp_path, monkeypatch):
    from generador_rfd import prototipo

    monkeypatch.setattr(prototipo.shutil, "which", lambda nombre: None)
    raiz = tmp_path / "varios"
    (raiz / "web").mkdir(parents=True)
    (raiz / "web" / "package.json").write_text(json.dumps({"scripts": {"dev": "next dev"}}), encoding="utf-8")
    (raiz / "node_modules" / "x").mkdir(parents=True)
    (raiz / "node_modules" / "x" / "package.json").write_text(json.dumps({"scripts": {"dev": "x"}}), encoding="utf-8")
    (c,) = prototipo.detectar_comandos(raiz)
    assert c["cwd"] == str(raiz / "web") and c["url_esperada"] == "http://localhost:3000"
    with pytest.raises(ValueError, match="carpeta"):
        prototipo.detectar_comandos(tmp_path / "no_existe")


def test_diseno_estatico_junta_el_css_real_y_omite_dependencias(tmp_path):
    from generador_rfd import prototipo

    raiz = _proyecto_web(tmp_path, **{
        "src/styles.css": ":root { --primario: #123456; }\nbody { font-family: Inter; }",
        "src/app.min.css": "x{}", "tailwind.config.js": "module.exports = { theme: { colors: { marca: '#ff0000' } } }",
        "node_modules/lib/estilo.css": "a{color:red}", "dist/salida.css": "a{color:blue}",
        "src/enorme.css": "a{color:green}\n" * 50_000,
    })
    texto = prototipo.diseno_estatico(raiz)
    assert "--primario: #123456" in texto and "marca: '#ff0000'" in texto
    assert "node_modules" not in texto and "salida.css" not in texto and "app.min.css" not in texto
    assert len(texto) < prototipo.MAX_DISENO + 2_000
    assert texto.index("tailwind.config.js") < texto.index("styles.css"), "primero la configuración de tokens"
    assert prototipo.diseno_estatico(tmp_path / "vacio") == ""


class _ProcesoFalso:
    """Sustituye a subprocess.Popen: nunca se inicia un proceso real."""

    creados: list = []

    def __init__(self, comando, salida: str = "", codigo=None, **opciones):
        import io

        self.comando, self.opciones, self.pid, self._codigo = comando, opciones, 4321, codigo
        self.stdout = io.StringIO(salida)
        _ProcesoFalso.creados.append(self)

    def poll(self):
        return self._codigo

    def terminate(self):
        self._codigo = -15

    kill = terminate

    def wait(self, timeout=None):
        return self._codigo


def _falso(monkeypatch, salida="", codigo=None):
    from generador_rfd import prototipo

    _ProcesoFalso.creados = []
    monkeypatch.setattr(prototipo.subprocess, "Popen", lambda comando, **o: _ProcesoFalso(comando, salida, codigo, **o))
    monkeypatch.setattr(prototipo, "_puerto_abierto", lambda host, puerto: False)
    return prototipo.Gestor()


def _hasta(condicion, limite=5.0):
    fin = time.time() + limite
    while not condicion():
        assert time.time() < fin, "la condición no se cumplió"
        time.sleep(0.01)


def test_levanta_sin_shell_y_detecta_la_url_de_la_salida(monkeypatch, tmp_path):
    gestor = _falso(monkeypatch, "\x1b[32m  VITE ready\x1b[0m\n  \x1b[1mLocal\x1b[0m: http://localhost:5173/\n")
    e = gestor.levantar("p1", ["bun", "run", "dev"], tmp_path)
    assert e["estado"] in ("iniciando", "activo")
    _hasta(lambda: gestor.estado("p1")["estado"] == "activo")
    e = gestor.estado("p1")
    (p,) = _ProcesoFalso.creados
    assert p.comando == ["bun", "run", "dev"] and not p.opciones.get("shell") and p.opciones["cwd"] == str(tmp_path)
    assert e["url"] == "http://localhost:5173" and "VITE ready" in "\n".join(e["log"]) and "\x1b" not in "\n".join(e["log"])


def test_sin_url_en_la_salida_prueba_el_puerto_esperado(monkeypatch, tmp_path):
    from generador_rfd import prototipo

    gestor = _falso(monkeypatch, "Serving HTTP on :: port 8765 (http://[::]:8765/)\n")
    gestor.levantar("p1", ["python", "-m", "http.server", "8765"], tmp_path, "http://127.0.0.1:8765")
    _hasta(lambda: gestor._procesos["p1"].lector_terminado)
    assert gestor.estado("p1")["estado"] == "iniciando"
    monkeypatch.setattr(prototipo, "_puerto_abierto", lambda host, puerto: (host, puerto) == ("127.0.0.1", 8765))
    e = gestor.estado("p1")
    assert e["estado"] == "activo" and e["url"] == "http://127.0.0.1:8765"


def test_el_log_esta_acotado(monkeypatch, tmp_path):
    from generador_rfd import prototipo

    gestor = _falso(monkeypatch, "".join(f"linea {i}\n" for i in range(prototipo.MAX_LINEAS_LOG + 50)))
    gestor.levantar("p1", ["x"], tmp_path)
    _hasta(lambda: gestor._procesos["p1"].lector_terminado)
    log = gestor.estado("p1")["log"]
    assert len(log) == prototipo.MAX_LINEAS_LOG and log[-1] == f"linea {prototipo.MAX_LINEAS_LOG + 49}"


def test_un_proceso_que_termina_queda_en_error_o_detenido(monkeypatch, tmp_path):
    gestor = _falso(monkeypatch, "Error: puerto ocupado\n", codigo=1)
    gestor.levantar("fallo", ["x"], tmp_path)
    _hasta(lambda: gestor._procesos["fallo"].lector_terminado)
    e = gestor.estado("fallo")
    assert e["estado"] == "error" and e["codigo"] == 1 and "puerto ocupado" in e["log"][-1]
    ok = _falso(monkeypatch, "", codigo=0)
    ok.levantar("fin", ["x"], tmp_path)
    assert ok.estado("fin")["estado"] == "detenido"
    assert ok.estado("nunca")["estado"] == "detenido"


def test_un_solo_proceso_por_prototipo_y_errores_claros(monkeypatch, tmp_path):
    from generador_rfd import prototipo

    gestor = _falso(monkeypatch)
    gestor.levantar("p1", ["x"], tmp_path)
    with pytest.raises(RuntimeError, match="Ya hay un proceso"):
        gestor.levantar("p1", ["y"], tmp_path)
    with pytest.raises(ValueError, match="comando"):
        gestor.levantar("p2", [], tmp_path)

    def sin_ejecutable(*a, **k):
        raise FileNotFoundError("no existe")

    monkeypatch.setattr(prototipo.subprocess, "Popen", sin_ejecutable)
    with pytest.raises(RuntimeError, match="No se pudo iniciar"):
        gestor.levantar("p3", ["inexistente"], tmp_path)
    assert gestor.estado("p3")["estado"] == "detenido"


@pytest.mark.parametrize("windows", [True, False])
def test_detener_mata_todo_el_arbol_de_procesos(monkeypatch, tmp_path, windows):
    from generador_rfd import prototipo

    gestor = _falso(monkeypatch)
    monkeypatch.setattr(prototipo, "WINDOWS", windows)
    llamadas = []
    monkeypatch.setattr(prototipo.subprocess, "run", lambda comando, **k: llamadas.append(("run", comando)))
    monkeypatch.setattr(prototipo.os, "killpg", lambda pgid, senal: llamadas.append(("killpg", pgid)), raising=False)
    monkeypatch.setattr(prototipo.os, "getpgid", lambda pid: pid, raising=False)
    gestor.levantar("p1", ["x"], tmp_path)
    assert gestor.detener("p1")["estado"] == "detenido"
    assert llamadas[0] == (("run", ["taskkill", "/PID", "4321", "/T", "/F"]) if windows else ("killpg", 4321))
    assert gestor.estado("p1")["estado"] == "detenido"
    gestor.levantar("p1", ["x"], tmp_path)  # tras detenerlo puede volver a levantarse


def test_detener_todos_y_url_ya_corriendo(monkeypatch, tmp_path):
    from generador_rfd import prototipo

    gestor = _falso(monkeypatch)
    monkeypatch.setattr(prototipo.subprocess, "run", lambda *a, **k: None)
    monkeypatch.setattr(prototipo.os, "killpg", lambda *a: None, raising=False)
    monkeypatch.setattr(prototipo.os, "getpgid", lambda pid: pid, raising=False)
    gestor.levantar("a", ["x"], tmp_path)
    gestor.levantar("b", ["x"], tmp_path)
    gestor.adoptar_url("c", "http://localhost:3000/")
    assert gestor.estado("c") == {**gestor.estado("c"), "estado": "activo", "url": "http://localhost:3000", "externo": True}
    with pytest.raises(ValueError, match="URL"):
        gestor.adoptar_url("d", "ftp://x")
    gestor.detener_todos()
    assert all(gestor.estado(i)["estado"] == "detenido" for i in "abc")


class _PaginaFalsa:
    def __init__(self, visitadas):
        self.visitadas = visitadas

    def goto(self, url, **k):
        self.visitadas.append(url)
        self.url = url

    def wait_for_timeout(self, ms):
        pass

    def screenshot(self, path, full_page):
        assert full_page is True
        Path(path).write_bytes(b"\x89PNG-falso")

    def evaluate(self, js):
        return {"html": "<html>" + "x" * 300_000 + "</html>", "css": f".a{{color:red}} /* {self.url} */",
                "diseno": {"fuentes": ["Inter"], "texto": "rgb(0, 0, 0)", "fondo": "rgb(255, 255, 255)", "acento": "rgb(1, 2, 3)",
                           "variables": {"--primario": "#123456"}}}


def _playwright_falso(monkeypatch, fallan=()):
    from contextlib import contextmanager

    from generador_rfd import prototipo

    visitadas, canales = [], []

    class Chromium:
        def launch(self, channel=None, **k):
            canales.append(channel)
            if channel in fallan:
                raise RuntimeError(f"{channel} no está instalado")
            return type("N", (), {"new_page": lambda s, **k: _PaginaFalsa(visitadas), "close": lambda s: None})()

    @contextmanager
    def sync_playwright():
        yield type("P", (), {"chromium": Chromium()})()

    monkeypatch.setattr(prototipo, "_sync_playwright", lambda: sync_playwright)
    return visitadas, canales


def test_captura_pantallas_css_y_resumen_de_diseno(monkeypatch, tmp_path):
    from generador_rfd import prototipo

    visitadas, canales = _playwright_falso(monkeypatch)
    destino = tmp_path / "captura"
    r = prototipo.capturar("http://localhost:5173/", ["/", "ordenes/nueva"], destino)
    assert visitadas == ["http://localhost:5173/", "http://localhost:5173/ordenes/nueva"] and canales == ["chrome"]
    assert [x["ruta"] for x in r] == ["/", "/ordenes/nueva"]
    assert (destino / "pantalla_1.png").read_bytes().startswith(b"\x89PNG")
    assert (destino / "pantalla_2.html").stat().st_size <= prototipo.MAX_HTML_PANTALLA + 100
    css = (destino / "estilos.css").read_text(encoding="utf-8")
    assert "localhost:5173/ */" in css and "ordenes/nueva" in css
    diseno = json.loads((destino / "diseno.json").read_text(encoding="utf-8"))
    assert diseno["variables"]["--primario"] == "#123456" and diseno["fuentes"] == ["Inter"]
    assert prototipo.pantallas_capturadas(destino) == r
    contexto = prototipo.contexto_captura(destino)
    assert "--primario" in contexto and "color:red" in contexto and "<html>" in contexto


def test_captura_cae_a_edge_y_explica_cuando_no_hay_navegador(monkeypatch, tmp_path):
    from generador_rfd import prototipo

    _, canales = _playwright_falso(monkeypatch, fallan=("chrome",))
    prototipo.capturar("http://localhost:3000", ["/"], tmp_path / "a")
    assert canales == ["chrome", "msedge"]
    _playwright_falso(monkeypatch, fallan=("chrome", "msedge"))
    with pytest.raises(RuntimeError, match="Chrome o Edge"):
        prototipo.capturar("http://localhost:3000", ["/"], tmp_path / "b")


def test_captura_sin_playwright_indica_como_instalarlo(monkeypatch, tmp_path):
    import builtins

    from generador_rfd import prototipo

    real = builtins.__import__

    def sin_playwright(nombre, *a, **k):
        if nombre.startswith("playwright"):
            raise ImportError(nombre)
        return real(nombre, *a, **k)

    monkeypatch.setattr(builtins, "__import__", sin_playwright)
    with pytest.raises(RuntimeError, match=r"pip install -e \"\.\[prototipo\]\""):
        prototipo.capturar("http://localhost:3000", ["/"], tmp_path / "c")


def test_captura_valida_la_url(monkeypatch, tmp_path):
    from generador_rfd import prototipo

    visitadas, _ = _playwright_falso(monkeypatch)
    for mala in ("file:///c:/secreto", "javascript:alert(1)", ""):
        with pytest.raises(ValueError, match="URL"):
            prototipo.capturar(mala, ["/"], tmp_path / "d")
    prototipo.capturar("http://localhost:3000", ["//otro.com/x", "/ok"], tmp_path / "d")
    assert visitadas == ["http://localhost:3000/otro.com/x", "http://localhost:3000/ok"]


def _html_proto(texto="Nueva pantalla"):
    return f'<!doctype html><html><head><style>.btn{{color:#123456}}</style></head><body><button class="btn">{texto}</button></body></html>'


def _modelo_falso(monkeypatch, respuestas):
    """Sustituye invocar(): devuelve las respuestas en orden y registra los prompts recibidos."""
    from generador_rfd import prototipo

    prompts = []

    def invocar_falso(*, proveedor, modelo, prompt, esquema, directorio, **k):
        prompts.append({"prompt": prompt, "esquema": esquema, "proveedor": proveedor, "modelo": modelo})
        return respuestas.pop(0)

    monkeypatch.setattr(prototipo, "invocar", invocar_falso)
    return prompts


def _nuevo_proto(tmp_path, rfd_id=None, **css):
    from generador_rfd import prototipo

    proyecto = _proyecto_web(tmp_path, **{"src/estilos.css": ":root{--marca:#abcdef}", **css})
    raiz = tmp_path / "prototipos"
    return prototipo.crear_prototipo(raiz, proyecto, rfd_id), raiz, proyecto


def test_turno_genera_la_primera_version_con_las_reglas_del_prompt(monkeypatch, tmp_path):
    from generador_rfd import prototipo

    id_, raiz, proyecto = _nuevo_proto(tmp_path)
    prompts = _modelo_falso(monkeypatch, [{"html": _html_proto(), "resumen": "Pantalla de ejemplo; falta el logotipo."}])
    datos_rfd = {"titulo": "Registro de órdenes", "objetivo": "Permitir registrar órdenes.", "requerimientos": [
        {"id": "RF-01", "nombre": "Registrar orden", "parrafos": ["El usuario registra."], "elementos": [{"titulo": "Campo cliente", "detalle": "x"}],
         "flujo": [{"numero": 1, "actor": "Usuario", "componente": "Formulario", "accion": "Completa los datos"}], "criterios_aceptacion": ["Se guarda"]}]}
    r = prototipo.turno(raiz / id_, "Muestra el formulario de registro", proveedor="claude", modelo="haiku", directorio=tmp_path, rfd=datos_rfd)
    assert r["version"] == 1 and "logotipo" in r["resumen"]
    (p,) = prompts
    assert p["esquema"]["required"] == ["html", "resumen"] and p["modelo"] == "haiku"
    for frase in ("autocontenido", "REUTILIZA", "--marca:#abcdef", "Nuevo", "Modificado", "sin CDN", "No introduzcas un nuevo lenguaje de diseño",
                  "Registrar orden", "Completa los datos", "Muestra el formulario de registro", "marcadores de posición"):
        assert frase in p["prompt"], frase
    assert "Prototipo actual" not in p["prompt"]
    assert (raiz / id_ / "v1.html").read_text(encoding="utf-8") == _html_proto()
    chat = prototipo.leer_chat(raiz / id_)
    assert [m["rol"] for m in chat] == ["usuario", "ia"] and chat[1]["version"] == 1 and chat[0]["texto"].startswith("Muestra")
    assert chat[0]["cuando"] and prototipo.versiones(raiz / id_) == [1]


def test_turno_itera_sobre_la_version_anterior_y_conserva_las_previas(monkeypatch, tmp_path):
    from generador_rfd import prototipo

    id_, raiz, _ = _nuevo_proto(tmp_path)
    prompts = _modelo_falso(monkeypatch, [{"html": _html_proto("Uno"), "resumen": "Primera."}, {"html": _html_proto("Dos"), "resumen": "Segunda."}])
    prototipo.turno(raiz / id_, "Primera pantalla", proveedor="claude", modelo="", directorio=tmp_path)
    r = prototipo.turno(raiz / id_, "Cambia el botón", proveedor="claude", modelo="", directorio=tmp_path)
    assert r["version"] == 2 and "Uno" in prompts[1]["prompt"] and "Prototipo actual" in prompts[1]["prompt"]
    assert "Primera pantalla" in prompts[1]["prompt"], "la conversación previa viaja en el prompt"
    assert "Uno" in (raiz / id_ / "v1.html").read_text(encoding="utf-8") and "Dos" in (raiz / id_ / "v2.html").read_text(encoding="utf-8")
    assert prototipo.versiones(raiz / id_) == [1, 2] and len(prototipo.leer_chat(raiz / id_)) == 4


def test_turno_usa_la_captura_si_existe_y_si_no_el_diseno_estatico(monkeypatch, tmp_path):
    from generador_rfd import prototipo

    id_, raiz, _ = _nuevo_proto(tmp_path)
    prompts = _modelo_falso(monkeypatch, [{"html": _html_proto(), "resumen": "a"}, {"html": _html_proto(), "resumen": "b"}])
    prototipo.turno(raiz / id_, "uno", proveedor="claude", modelo="", directorio=tmp_path)
    assert "--marca:#abcdef" in prompts[0]["prompt"] and "CSS completo del proyecto" not in prompts[0]["prompt"]
    _playwright_falso(monkeypatch)
    prototipo.capturar("http://localhost:5173", ["/"], raiz / id_ / "captura")
    prototipo.turno(raiz / id_, "dos", proveedor="claude", modelo="", directorio=tmp_path)
    assert "CSS completo del proyecto" in prompts[1]["prompt"] and "--primario" in prompts[1]["prompt"]


def test_turno_rechaza_respuestas_invalidas_y_avisa_de_recursos_externos(monkeypatch, tmp_path):
    from generador_rfd import prototipo

    id_, raiz, _ = _nuevo_proto(tmp_path)
    externo = '<!doctype html><html><head><link rel="stylesheet" href="https://cdn.ejemplo.com/x.css"></head><body><a href="https://ok.com">enlace</a></body></html>'
    _modelo_falso(monkeypatch, [{"html": "no es html", "resumen": "x"}, {"html": "", "resumen": "x"}, {"html": externo, "resumen": "Con estilos."}])
    for _ in range(2):
        with pytest.raises(ValueError, match="HTML"):
            prototipo.turno(raiz / id_, "m", proveedor="claude", modelo="", directorio=tmp_path)
    assert prototipo.leer_chat(raiz / id_) == [] and prototipo.versiones(raiz / id_) == []
    r = prototipo.turno(raiz / id_, "m", proveedor="claude", modelo="", directorio=tmp_path)
    assert "cdn.ejemplo.com" in r["resumen"] and "ok.com" not in r["resumen"] and r["resumen"].startswith("Con estilos.")


def test_crear_prototipo_valida_la_ruta_y_los_ids(tmp_path):
    from generador_rfd import prototipo

    with pytest.raises(ValueError, match="carpeta"):
        prototipo.crear_prototipo(tmp_path / "p", tmp_path / "no_existe")
    id_ = prototipo.crear_prototipo(tmp_path / "p", tmp_path, 7)
    assert prototipo.leer_meta(tmp_path / "p" / id_)["rfd_id"] == 7
    assert prototipo.carpeta_de(tmp_path / "p", id_) == tmp_path / "p" / id_
    for malo in ("../x", "..", "a/b", "a" * 80, "no_existe_ok"):
        with pytest.raises(KeyError):
            prototipo.carpeta_de(tmp_path / "p", malo)


@pytest.fixture
def proto(api, tmp_path, monkeypatch):
    """API con un gestor de procesos y un Playwright falsos: ningún test inicia procesos ni navegadores."""
    from generador_rfd import prototipo

    gestor = _falso(monkeypatch)
    monkeypatch.setattr(prototipo, "GESTOR", gestor)
    monkeypatch.setattr(prototipo.subprocess, "run", lambda *a, **k: None)
    monkeypatch.setattr(prototipo.os, "killpg", lambda *a: None, raising=False)
    monkeypatch.setattr(prototipo.os, "getpgid", lambda pid: pid, raising=False)
    monkeypatch.setattr(prototipo.shutil, "which", lambda nombre: None)
    api.proyecto = _proyecto_web(tmp_path, **{"src/estilos.css": ":root{--marca:#abcdef}"})
    api.gestor = gestor
    return api


def _crear_proto(cliente, rfd_id=None) -> str:
    r = cliente.post("/api/prototipos", json={"ruta": str(cliente.proyecto), "rfd_id": rfd_id})
    assert r.status_code == 200, r.text
    return r.json()["id"]


def test_api_detecta_comandos_y_valida_la_ruta(proto):
    r = proto.get("/api/prototipos/detectar", params={"ruta": str(proto.proyecto)})
    assert r.status_code == 200 and r.json()["comandos"][0]["comando"] == ["npm", "run", "dev"]
    assert proto.get("/api/prototipos/detectar", params={"ruta": str(proto.proyecto / "nada")}).status_code == 400
    assert proto.post("/api/prototipos", json={"ruta": str(proto.proyecto / "nada")}).status_code == 400
    assert proto.post("/api/prototipos", json={"ruta": str(proto.proyecto), "rfd_id": 999}).status_code == 404


def test_api_crea_y_consulta_un_prototipo(proto):
    id_ = _crear_proto(proto)
    d = proto.get(f"/api/prototipos/{id_}").json()
    assert d["ruta"] == str(proto.proyecto.resolve()) and d["rfd"] is None and d["chat"] == [] and d["versiones"] == []
    assert d["proceso"]["estado"] == "detenido" and d["captura"]["pantallas"] == []
    (fila,) = proto.get("/api/prototipos").json()
    assert fila["id"] == id_ and fila["ruta"] == d["ruta"] and fila["versiones"] == 0 and fila["rfd"] is None
    for malo in ("0000000000", "zzzz", "..%2F..%2Fx"):
        assert proto.get(f"/api/prototipos/{malo}").status_code == 404


def test_api_levanta_detiene_y_exige_un_comando_valido(proto):
    id_ = _crear_proto(proto)
    r = proto.post(f"/api/prototipos/{id_}/levantar", json={"comando": ["npm", "run", "dev"], "url_esperada": "http://localhost:5173"})
    assert r.status_code == 200, r.text
    (p,) = _ProcesoFalso.creados
    assert p.comando == ["npm", "run", "dev"] and p.opciones["cwd"] == str(proto.proyecto.resolve())
    assert proto.post(f"/api/prototipos/{id_}/levantar", json={"comando": ["npm", "start"]}).status_code == 409
    assert proto.post(f"/api/prototipos/{id_}/detener").json()["estado"] == "detenido"
    otro = proto.proyecto.parent
    assert proto.post(f"/api/prototipos/{id_}/levantar", json={"comando": ["x"], "cwd": str(otro)}).status_code == 400
    assert proto.post(f"/api/prototipos/{id_}/levantar", json={"comando": []}).status_code == 400
    assert proto.post(f"/api/prototipos/{id_}/levantar", json={}).status_code == 400
    # «Ya está corriendo»: sólo se recuerda la URL, sin iniciar nada.
    antes = len(_ProcesoFalso.creados)
    d = proto.post(f"/api/prototipos/{id_}/levantar", json={"url": "http://localhost:3000/"}).json()
    assert d["estado"] == "activo" and d["url"] == "http://localhost:3000" and len(_ProcesoFalso.creados) == antes
    assert proto.post(f"/api/prototipos/{id_}/levantar", json={"url": "ftp://x"}).status_code == 400


def test_api_captura_en_segundo_plano_y_sirve_solo_su_carpeta(proto, monkeypatch):
    id_ = _crear_proto(proto)
    assert proto.post(f"/api/prototipos/{id_}/capturar", json={"rutas": ["/"]}).status_code == 409, "sin proyecto en ejecución ni URL"
    _playwright_falso(monkeypatch)
    proto.post(f"/api/prototipos/{id_}/levantar", json={"url": "http://localhost:5173"})
    r = proto.post(f"/api/prototipos/{id_}/capturar", json={"rutas": ["/", "/ordenes"]})
    assert r.status_code == 200, r.text
    t = _esperar(proto, r.json()["id"])
    assert t["estado"] == "listo", t.get("error")
    assert [x["ruta"] for x in t["resultado"]["pantallas"]] == ["/", "/ordenes"]
    d = proto.get(f"/api/prototipos/{id_}").json()
    assert [x["imagen"] for x in d["captura"]["pantallas"]] == ["pantalla_1.png", "pantalla_2.png"]
    png = proto.get(f"/api/prototipos/{id_}/captura/pantalla_1.png")
    assert png.status_code == 200 and png.content.startswith(b"\x89PNG") and png.headers["content-type"] == "image/png"
    html = proto.get(f"/api/prototipos/{id_}/captura/pantalla_1.html")
    assert html.headers["content-type"].startswith("text/plain")
    for malo in ("..%2Fmeta.json", "..%5Cmeta.json", "inexistente.png", "%2e%2e%2fmeta.json", "meta.json"):
        assert proto.get(f"/api/prototipos/{id_}/captura/{malo}").status_code == 404
    assert proto.get(f"/api/prototipos/{id_}/captura/captura.json").status_code == 200


def test_api_captura_sin_playwright_falla_con_un_mensaje_claro(proto, monkeypatch):
    from generador_rfd import prototipo

    def sin_playwright():
        raise RuntimeError(prototipo.INSTALAR_PLAYWRIGHT)

    monkeypatch.setattr(prototipo, "_sync_playwright", sin_playwright)
    id_ = _crear_proto(proto)
    proto.post(f"/api/prototipos/{id_}/levantar", json={"url": "http://localhost:5173"})
    t = _esperar(proto, proto.post(f"/api/prototipos/{id_}/capturar", json={"rutas": ["/"]}).json()["id"])
    assert t["estado"] == "error" and 'pip install -e ".[prototipo]"' in t["error"]


def test_api_chat_genera_versiones_con_el_rfd_vinculado(proto, tmp_path, monkeypatch):
    rfd_id, _ = _importar_por_api(proto, tmp_path, monkeypatch)
    id_ = _crear_proto(proto, rfd_id)
    prompts = _modelo_falso(monkeypatch, [{"html": _html_proto("Uno"), "resumen": "Primera."}, {"html": _html_proto("Dos"), "resumen": "Segunda."}])
    assert proto.post(f"/api/prototipos/{id_}/mensaje", data={"texto": "  ", "proveedor": "claude"}).status_code == 400
    r = proto.post(f"/api/prototipos/{id_}/mensaje", data={"texto": "Muestra el registro", "proveedor": "claude", "modelo": "haiku"})
    assert r.status_code == 200, r.text
    t = _esperar(proto, r.json()["id"])
    assert t["estado"] == "listo", t.get("error")
    assert t["resultado"]["version"] == 1 and t["resultado"]["resumen"] == "Primera." and t["tipo"] == "prototipo"
    assert prompts[0]["modelo"] == "haiku" and servidor.almacen().obtener(rfd_id)["datos"]["titulo"] in prompts[0]["prompt"]
    t2 = _esperar(proto, proto.post(f"/api/prototipos/{id_}/mensaje", data={"texto": "Cambia", "proveedor": "claude"}).json()["id"])
    assert t2["resultado"]["version"] == 2
    d = proto.get(f"/api/prototipos/{id_}").json()
    assert d["versiones"] == [1, 2] and [m["rol"] for m in d["chat"]] == ["usuario", "ia", "usuario", "ia"]
    assert d["rfd"]["id"] == rfd_id and d["rfd"]["codigo"] == "RFD-9"
    v = proto.get(f"/api/prototipos/{id_}/v/1")
    assert v.status_code == 200 and "Uno" in v.text and v.headers["content-type"].startswith("text/html")
    csp = v.headers["content-security-policy"]
    assert csp.startswith("sandbox allow-scripts") and "default-src 'none'" in csp and "script-src 'unsafe-inline'" in csp
    assert proto.get(f"/api/prototipos/{id_}/v/9").status_code == 404


def test_api_rechaza_un_segundo_trabajo_simultaneo_del_mismo_prototipo(proto):
    id_ = _crear_proto(proto)
    servidor._proto_ocupados.add(id_)
    try:
        assert proto.post(f"/api/prototipos/{id_}/mensaje", data={"texto": "hola", "proveedor": "claude"}).status_code == 409
    finally:
        servidor._proto_ocupados.discard(id_)


def test_api_adjunta_el_prototipo_junto_al_word_del_rfd(proto, tmp_path, monkeypatch):
    rfd_id, _ = _importar_por_api(proto, tmp_path, monkeypatch)
    _modelo_falso(monkeypatch, [{"html": _html_proto("Uno"), "resumen": "a"}, {"html": _html_proto("Dos"), "resumen": "b"}])
    id_ = _crear_proto(proto, rfd_id)
    assert proto.post(f"/api/prototipos/{id_}/adjuntar-rfd", json={}).status_code == 409, "aún no hay versiones"
    for texto in ("uno", "dos"):
        _esperar(proto, proto.post(f"/api/prototipos/{id_}/mensaje", data={"texto": texto, "proveedor": "claude"}).json()["id"])
    assert proto.get(f"/api/rfds/{rfd_id}").json()["archivos"]["prototipo"] is False
    r = proto.post(f"/api/prototipos/{id_}/adjuntar-rfd", json={"version": 1})
    assert r.status_code == 200, r.text
    destino = Path(servidor.almacen().obtener(rfd_id)["ruta_docx"]).parent / "RFD_RFD-9_prototipo.html"
    assert destino.is_file() and "Uno" in destino.read_text(encoding="utf-8")
    assert proto.get(f"/api/rfds/{rfd_id}").json()["archivos"]["prototipo"] is True
    assert "Uno" in proto.get(f"/api/rfds/{rfd_id}/archivo/prototipo").text
    assert proto.post(f"/api/prototipos/{id_}/adjuntar-rfd", json={"version": 7}).status_code == 404
    sin_rfd = _crear_proto(proto)
    assert proto.post(f"/api/prototipos/{sin_rfd}/adjuntar-rfd", json={}).status_code == 400


def test_al_cerrar_la_aplicacion_se_detienen_los_procesos(proto):
    from fastapi.testclient import TestClient

    id_ = _crear_proto(proto)
    proto.post(f"/api/prototipos/{id_}/levantar", json={"comando": ["npm", "run", "dev"]})
    assert proto.gestor.estado(id_)["estado"] != "detenido"
    with TestClient(servidor.app):
        pass
    assert proto.gestor.estado(id_)["estado"] == "detenido"


def test_un_comando_escrito_como_texto_se_parte_en_argumentos(proto):
    from generador_rfd import prototipo

    assert prototipo.partir_comando('npm run dev -- --port "5174"') == ["npm", "run", "dev", "--", "--port", "5174"]
    assert prototipo.partir_comando(r'"C:\Program Files\nodejs\npm.cmd" start') == [r"C:\Program Files\nodejs\npm.cmd", "start"]
    id_ = _crear_proto(proto)
    assert proto.post(f"/api/prototipos/{id_}/levantar", json={"comando": "npm run dev"}).status_code == 200
    assert _ProcesoFalso.creados[-1].comando == ["npm", "run", "dev"]
